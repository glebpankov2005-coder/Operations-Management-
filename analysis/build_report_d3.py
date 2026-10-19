"""
Build the Deliverable 3 (Performance overview: KPIs + CAPEX/OPEX) Word document.
Reads outputs/reports/phase1_metrics.json (performance_cost, dock_mhe, layout, picking, capacity_plan sections),
figures and data/assumptions.json; writes deliverables/Deliverable_3_Performance_Overview.docx.
Run performance_cost.py first.
"""
import os, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..")
REPORTS = os.path.join(BASE, "outputs", "reports")
FIGS = os.path.join(BASE, "outputs", "figures")
DATA = os.path.join(BASE, "data")
M = json.load(open(os.path.join(REPORTS, "phase1_metrics.json"), encoding="utf-8"))
ASSUM = json.load(open(os.path.join(DATA, "assumptions.json"), encoding="utf-8"))

RED = RGBColor(0xD4, 0x00, 0x11); DARK = RGBColor(0x22, 0x22, 0x22); GREY = RGBColor(0x66, 0x66, 0x66)
doc = Document()
normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = DARK
for lvl, sz in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    st = doc.styles[lvl]; st.font.name = "Calibri"; st.font.size = Pt(sz)
    st.font.color.rgb = RED if lvl == "Heading 1" else DARK; st.font.bold = True


def para(t, size=10.5, color=DARK, bold=False, italic=False, align=None, sa=6):
    """Paragraph; '**text**' segments are bold (markers are not printed)."""
    p = doc.add_paragraph()
    for i, seg in enumerate(t.split("**")):
        if not seg:
            continue
        r = p.add_run(seg)
        r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold or (i % 2 == 1); r.italic = italic
    if align: p.alignment = align
    if t.rstrip().endswith(":"): p.paragraph_format.keep_with_next = True   # lead-in stays with its table
    p.paragraph_format.space_after = Pt(sa); return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        for i, seg in enumerate(it.split("**")):
            r = p.add_run(seg); r.bold = (i % 2 == 1); r.font.size = Pt(10.5)


def table(headers, rows, widths=None, arr=1, bold_last=False):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = ""; r = t.rows[0].cells[i].paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(9.5)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(9.5)
            r.bold = bold_last and ri == len(rows) - 1
            if i >= arr: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, w in enumerate(widths):
            grid.gridCol_lst[i].w = Cm(w)                      # grid widths: honoured by Word and LibreOffice
            for row in t.rows: row.cells[i].width = Cm(w)
    for ri, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if ri < len(t.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs: p.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t


def figure(name, cap, w=15.5):
    p = os.path.join(FIGS, name)
    if not os.path.exists(p):
        raise FileNotFoundError(p)
    doc.add_picture(p, width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(cap, size=8.5, color=GREY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)


def footer():
    f = doc.sections[0].footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = f.add_run("Project OTM — Assa Abloy @ DHL Bemmel  ·  Deliverable 3 (Performance overview)  ·  Page")
    run.font.size = Pt(8); run.font.color.rgb = GREY
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); run._r.addnext(fld)


def eur(v, k=False):
    return f"€{v / 1000:,.0f}k" if k else f"€{v:,.0f}"


PC = M["performance_cost"]; DM = M["dock_mhe"]; LAY = M["layout"]; PK = M["picking"]; CP = M["capacity_plan"]
Q = PC["quantities"]; LT = PC["leadtime"]; SORT = PC["sorting"]; STO = PC["storage"]; LZ = DM["loose_zone"]
fleet = {f["Model"].replace("Jungheinrich ", ""): f for f in DM["fleet"]}
seg = {r["segment"]: r for r in PK["segments"]}
kp = {k["KPI"]: k for k in PC["kpis"]}
LC = DM["load_checks"]
PTL_COST = sum(r["Total EUR"] for r in PC["capex_table"] if r["Item"].startswith("Put-to-light"))

# ================================================================ COVER
SP = PK["small_parts"]; OA = DM["outbound_area"]; PA = DM["packing_area"]; HZ = M["design_evaluation"]["hazmat"]
LG = LC["long"]; SO = {o["Option"]: o for o in PC["small_item_options"]}
SO_SHELF = [o for o in PC["small_item_options"] if o["Option"].startswith("Bin")][0]
SO_LIFT = [o for o in PC["small_item_options"] if o["Option"].startswith("Kardex")][0]
SO_RACK = [o for o in PC["small_item_options"] if o["Option"].startswith("Keep")][0]
RATE = PC["cost_register"]["labour_rate"]
doc.add_paragraph().paragraph_format.space_before = Pt(55)
para("PROJECT OTM", size=13, color=RED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
para("Performance Overview — KPIs, CAPEX & OPEX", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
para("Assa Abloy operation · DHL Supply Chain, Bemmel (NL_0032)", size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=26)
table(["Field", "Detail"], [
    ["Deliverable", "3 — Performance overview including financial analysis"],
    ["Due date", "21-10-2026"],
    ["Concept", "Option B — single-deep racks, fast-pick area, small-parts shelving, zone + batch picking, wave release"],
    ["Envelope", f"≤ 7,000 m² · ≤ 12.2 m (design uses {LAY['total_area_m2']:,} m²)"],
    ["Team", "Group 3 — Gleb, Rachel Selvarathnam, Arnau, Vidhi"],
    ["Version", "v2 — includes the feedback of the last session (space checks, lengths, small parts, own KPIs, €20/h, hazmat)"],
    ["Supporting files", "Deliverable_3_Calculations.xlsx (all cost formulas) · Deliverables 1 and 2 (updated) · analysis/*.py"],
], widths=[4.5, 11], arr=9)
para("Expected performance and cost of the Option B design. Volumes are measured in the Assa Abloy Q2-2025 data; "
     "prices are public benchmarks or labelled assumptions (Section 7), never quotes. Every number can be rebuilt with "
     "analysis/performance_cost.py and checked in the calculation workbook.", size=9, color=GREY, italic=True)
doc.add_page_break()

doc.add_heading("Contents", level=1)
_p = doc.add_paragraph()
_fld = OxmlElement("w:fldSimple"); _fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
_r = OxmlElement("w:r"); _t = OxmlElement("w:t")
_t.text = "(In Word: right-click → Update Field to build the table of contents.)"
_r.append(_t); _fld.append(_r); _p._p.append(_fld)
doc.add_page_break()

# ================================================================ SUMMARY
doc.add_heading("Summary", level=1)
para("What the design delivers, what it costs, and what to watch:", sa=4)
table(["", "Result"], [
    ["Picking effort", f"{PC['pick_today_staffed_h']:.0f} → {PC['pick_future_h']:.0f} picker-hours a day (batching, fast-pick area, wave release)"],
    ["Order lead time", f"orders out within 24 h: {LT['today_within_24h_pct']:.0f}% → {LT['future_within_24h_pct']:.0f}%; "
                        f"slowest 10% wait {LT['today_p90_h']:.0f} h → {LT['future_p90_h']:.0f} h"],
    ["Trailers at the doors", f"{DM['outbound_doors']['current_drop_trailer_concurrent_p90']:.0f} parked at once → max "
                              f"{DM['departure_plan']['plan_max_doors_busy']} of {DM['outbound_doors']['doors_recommended']} doors"],
    ["Storage", f"{LAY['positions_total']:,} pallet positions + {SP['shelf_bays']} small-parts shelving bays; racks "
                f"{STO['rack_need_peak'] / LAY['positions_total']:.0%} full at the peak"],
    ["People and trucks", f"{STO['fte_avg']:.1f} FTE on an average day, {STO['fte_peak']:.1f} on the peak day; "
                          f"{Q['etv']} ETV 216i, {Q['eje']} EJE 120i, {Q['ere']} ERE 225, {Q['trolleys']} trolleys"],
    ["One-off investment (CAPEX)", f"{eur(PC['capex_total'], True)} all new; {eur(PC['capex_reuse_racking'], True)} if today's racking is re-used; "
                                   f"of which {eur(PC['capex_process'], True)} for the process changes"],
    ["Running cost (OPEX)", f"{eur(PC['opex_total'], True)} a year at €{RATE['value']:.0f}/h; {eur(PC['annual_cost'], True)} incl. depreciation"],
    ["Cost per unit", f"€{PC['unit_cost']['per_order_line']:.2f} per order line · €{PC['unit_cost']['per_order']:.2f} per order"],
    ["Saving and payback", f"~{eur(PC['saving_year'], True)} a year in picking labour → the process changes pay back in "
                           f"{PC['payback_process_years']:.1f} years"],
], widths=[4.4, 11.6], arr=9)
bullets([
    "**The labour gain comes from how we work, not from steel.** Batching, a fast-pick area and wave release need "
    f"{eur(PC['capex_process'], True)} and save ~{eur(PC['saving_year'], True)} a year. Racks are needed to hold the stock, "
    "so they are a capacity cost, not a saving.",
    f"**Small parts are stored efficiently:** {SP['skus']:,} slow SKUs with a few litres of stock each move from pallet "
    f"locations to bins next to the put-wall. That frees {SP['freed_rack_locations']:,} pallet locations. A Kardex lift was "
    f"compared: denser, but ~{eur(SO_LIFT['Yearly cost EUR'] - SO_SHELF['Yearly cost EUR'], True)} a year more and busy only a few % of the day at our volume.",
    f"**Space checks passed:** {LZ['share_fit_zone_pct']:.0f}% of containers fit the container zone, {LZ['share_fit_with_r2_pct']:.0f}% when door R2 "
    f"is lent; outbound staging uses at most {OA['need_m2_max'] / OA['planned_m2']:.0%} of its area; packing needs {PA['need_m2']:.0f} of 450 m².",
    f"**Watch the stock growth.** Occupied locations grew {STO['growth_per_quarter_pct']:.0f}% a quarter between April and June; at "
    f"that rate the racks are full in ~{STO['months_headroom']:.0f} months. Clearing dead stock (21% of loads > 1 year) is the main lever.",
])
doc.add_page_break()

# ================================================================ 1 WHAT CHANGED
doc.add_heading("1. What changed since Deliverable 2", level=1)
para("Two feedback rounds and a final check of the design led to these changes. Deliverables 1 and 2 were regenerated so "
     "all three documents use the same numbers.", sa=4)
table(["Topic", "Before", "Now", "Why"], [
    ["Racking", "Double-deep, telescopic forks", f"Single-deep, back-to-back ({LAY['positions_total']:,} positions)",
     "Team decision: no two-deep racks; every pallet directly reachable"],
    ["Long goods", "Cantilever for items up to 11 m", f"Long-goods module: {LAY['positions_long']} positions for pallets up to 2.05 m, 3.2 m aisle",
     "Feedback: nothing longer than a door (~2 m); the 11 m item is a data error"],
    ["Small parts", "One pallet location per SKU", f"Bin shelving for {SP['skus']:,} slow small SKUs next to the put-wall",
     "Feedback: store small items efficiently (Kardex compared)"],
    ["Loose containers", "Unpacked at door R1", f"Own {LZ['area_planned_m2']} m² zone; big containers borrow door R2",
     "Feedback: is there enough space to unpack?"],
    ["Reach / pallet trucks", "3 ETV 216i / 4 EJE 120i", f"{Q['etv']} ETV 216i / {Q['eje']} EJE 120i", "Peak need covered; lease includes repair"],
    ["Labour cost", "€28 per hour", f"€{RATE['value']:.0f} per hour (€{RATE['high']:.0f} as high case)", "Feedback: €20/h for an operator"],
    ["KPIs", "W3 list", "Our own design KPIs + the W3 ones", "Feedback: come up with our own KPIs"],
    ["Hazardous goods", "Class unknown, PGS 15", "Recounted, classified by customs code → mostly PGS 37-2", "Feedback: recheck counts and classes"],
], widths=[2.8, 3.6, 4.6, 5.0], arr=9)
figure("layout_orderflow.png", "Fig 1. Updated floor plan: container zone (C) with R1, long-goods module, the packing area with the "
       "put-to-light wall, and the main order flow (1–6).", 16.5)

# ================================================================ 2 KPIs
doc.add_heading("2. Our KPIs — expected performance of the design", level=1)
para("We made our own KPIs: each one measures something our design promises to do (fill the racks well, keep pickers "
     "in the fast-pick area, walk little, unpack containers inside their zone, ship within a day, …). The four W3 lecture "
     "KPIs (dock-to-stock, put-away time, pick rate, on-time shipment, S017) are included too. For every KPI we give how "
     "to calculate it, today's value where the data shows it, the value we expect from the design, and a target.", sa=4)
for grp in ["Space", "Inbound", "Picking", "Service", "Docks", "People", "Machinery", "Cost", "Safety"]:
    rows = [[k["KPI"] + (" ★" if k["Type"] == "own" else ""), k["How it is measured"], k["Today (data)"], k["Future (design)"], k["Target / benchmark"]]
            for k in PC["kpis"] if k["Group"] == grp]
    if not rows:
        continue
    para(f"**{grp}**", sa=2).paragraph_format.keep_with_next = True
    table(["KPI", "How we calculate it", "Today (data)", "Expected (design)", "Target"], rows,
          widths=[3.0, 3.6, 3.0, 3.8, 2.6], arr=9)
para("★ = our own design KPI; the others are the W3 lecture KPIs.", size=9, color=GREY, italic=True)

doc.add_heading("2.1 Order lead time", level=2)
figure("d3_leadtime.png", "Fig 2. Share of orders shipped within X hours — today (measured) vs future (wave plan, modelled).", 15)
bullets([
    f"**Today (measured, {LT['orders']:,} orders):** median {LT['today_median_h']:.0f} h, average {LT['today_mean_h']:.0f} h, "
    f"slowest 10% {LT['today_p90_h']:.0f} h or more. Only {LT['today_within_24h_pct']:.0f}% of orders leave within 24 hours.",
    "**Future (modelled on the same orders):** an order received before 12:00 is picked in wave 3 and leaves on the same "
    "afternoon's trailer of its carrier; a later order is picked in the first wave next morning and leaves on that day's "
    "first slot of its carrier. Weekends are skipped.",
    f"**Result:** {LT['future_within_24h_pct']:.0f}% within 24 hours and the slowest 10% at {LT['future_p90_h']:.0f} h. "
    f"The median moves less ({LT['today_median_h']:.0f} → {LT['future_median_h']:.0f} h) because "
    f"{DM['order_release']['orders_arrive_15_17_pct']:.0f}% of orders arrive between 15:00 and 17:00 and still leave the next day — "
    "the gain is that no order waits for a second or third day.",
])
doc.add_heading("2.2 Picking effort — where the saving comes from", level=2)
figure("d3_levers.png", "Fig 3. Picker-hours per average day: today's way of working (modelled) vs the design.", 15)
lv = PC["levers"]
table(["Lever", "Today h/day", "Future h/day", "Saved h/day", "Saved per year"],
      [[r["Lever"], f"{r['Today h/day']:.1f}", f"{r['Future h/day']:.1f}", f"{r['Saved h/day']:.1f}", eur(r["Saved EUR/yr"])] for r in lv],
      widths=[6.6, 2.2, 2.2, 2.2, 2.8])
para("Today is modelled with the same pick-time model and rates as the design (D2 §6.2), but the way the data says the "
     "site works now: one order per tour, stock not placed by demand (the WMS class agrees with demand for only 48.6% of "
     f"SKUs), and all orders released in the evening so the morning starts at {DM['order_release']['pick_peak_over_mean_active_hour']}× "
     "the average workload. It is a model, not a measurement — DHL's real pick hours would confirm it (A006).", size=9.5, color=GREY)

doc.add_heading("2.3 Storage space", level=2)
bullets([
    f"**Positions:** {LAY['positions_total']:,} (incl. {LAY['positions_long']} long-goods). Moving small parts to shelving frees "
    f"{SP['freed_rack_locations']:,} locations, so the peak need is ~{STO['rack_need_peak']:,} "
    f"({STO['rack_need_peak'] / LAY['positions_total']:.0%} full).",
    f"**Growth:** occupied locations rose from 7,715 (April) to 8,512 (June), +{STO['growth_per_quarter_pct']:.0f}% per quarter. "
    f"At that speed the racks are full in ~{STO['months_headroom']:.0f} months (~{STO['months_headroom_12']:.0f} with a 12th rack module, "
    f"+{LAY['positions_12th_module']} positions, which fits in the core). Three months of data cannot tell whether this is a season "
    "or a trend, so DHL should track the KPI monthly.",
    "**What we do about it:** (1) clear dead stock — 21% of loads are older than one year (D1 Fig 2b); (2) keep the 12th "
    "module as the first expansion; (3) the remaining ~270 m² of the 7,000 m² envelope stays free.",
])
doc.add_heading("2.4 Stress test — 150%, 200% and 250% volume", level=2)
para("Promised in D2 §3.1: what happens when volume grows? Workload scales with volume; trucks and doors are checked "
     "against the same rules as in D2.", sa=4)
figure("d3_stress.png", "Fig 4. Direct FTE and trucks needed on the peak day at higher volumes.", 13.5)
st_ = {s_["Volume"]: s_ for s_ in PC["stress"]}
table(["Volume", "Direct FTE (peak day)", "ETV 216i", "EJE 120i", "ERE 225", "Outbound doors", "Pickers in busiest wave"],
      [[s["Volume"], f"{s['Direct FTE (peak day)']:.1f}", s["ETV 216i needed"], s["EJE 120i needed"], s["ERE 225 needed"],
        s["Outbound doors (busiest hour)"], s["Pickers in the busiest wave"]] for s in PC["stress"]],
      widths=[2.0, 2.8, 1.8, 1.8, 1.8, 2.4, 3.0])
bullets([
    f"**Up to 150%** the design runs as planned: {st_['150%']['ETV 216i needed']} reach trucks, {st_['150%']['EJE 120i needed']} pallet trucks "
    f"and {st_['150%']['Outbound doors (busiest hour)']} shipping doors in the busiest hour.",
    "**At 200% and more** the 5 shipping doors are not enough in the busiest hour; the departure plan would need more spread "
    "(more slots before 12:00) or a 6th door — and storage would need the 12th module and a dead-stock clean-up first.",
    f"**Pickers per wave stay low** (≤ {st_['250%']['Pickers in the busiest wave']} in the busiest wave at 250%), so aisle congestion is not the limit.",
])

# ================================================================ 3 SORTING + SMALL PARTS
doc.add_heading("3. Sorting and small parts — what, how and when", level=1)
para("Sorting means putting goods into the right group for the next step. In our design it happens in four places. "
     "Each one is simple, uses a scan to check the item, and happens at a fixed moment in the day.", sa=4)
table(["Where", "What is sorted", "How", "When", "Equipment"], [
    ["1. Container zone (door R1)", "Cartons from a container → one pallet per SKU",
     f"Cartons roll out on the conveyor; the worker scans each carton and puts it on the build spot of its SKU (a container has "
     f"{SORT['container_skus_median']:.0f} SKUs on average, so {SORT['build_spots']} build spots are enough). Full pallet → wrap, label, putaway.",
     "Morning, when the container arrives (~3 h)", "Roller conveyor, RF scanner, stretch wrapper"],
    ["2. Put-to-light wall (packing)", "Picked items → customer orders",
     f"Pickers collect the items of many small orders in one trolley tour (~{SORT['lines_per_tour']:.0f} lines); small parts come "
     "straight from the shelving next to the wall. The packer scans an item, a light shows the order's slot, the item goes in and "
     "the button is pressed. A full slot lights up → pack.",
     f"During each wave (3 a day, ~{SORT['small_orders_wave_avg']:.0f} small orders per wave)", f"Wall with {SORT['ptl_slots']} slots"],
    ["3. Carrier lanes (outbound)", "Packed parcels and pallets → trailer",
     f"The label says the carrier and departure slot. Parcels go into roll cages per carrier, pallets into the staging lane of "
     f"their trailer ({SORT['staging_lanes']} trailer slots a day).", "All afternoon, until each trailer's departure",
     f"{SORT['roll_cages']} roll cages, floor lanes"],
    ["4. Returns", "Returned goods → restock / rework / scrap", "Inspect and grade, then each group goes to its own cage.",
     "When a return arrives", "Returns area, RF scanner"],
], widths=[2.8, 2.8, 5.6, 2.6, 2.2], arr=9)
figure("d3_sorting_timeline.png", "Fig 5. When the sorting points are busy during a normal day.", 16)
bullets([
    f"**Why a put-to-light wall and not sorting by hand:** {seg['small']['pct_orders']:.0f}% of orders are small (median "
    f"{PK['order_volume']['median_l']:.0f} litres). Picking them one by one would mean ~{SORT['small_orders_day']:.0f} walks a day; "
    "picking them together and sorting at the wall is about twice as fast (D2 §6.2).",
    f"**Why not an automatic sorter:** ~{SORT['small_orders_day']:.0f} small orders a day is far too little for a sorter "
    f"machine to pay off; the wall costs about {eur(PTL_COST, True)} (Section 5).",
    "**Large orders are not sorted:** each one is picked straight onto its own pallet, so it goes directly to packing.",
])
PWD = PK["putwall"]; TS = {int(k): v for k, v in PK["small_order_totes_share"].items()}
doc.add_heading("3.1 The put-to-light wall — where it is and how it works", level=2)
para("The put-to-light wall is where the small orders are sorted. Pickers do not pick order by order: they pick many "
     "small orders at once into a few totes (a batch), and the wall splits that batch into customer orders. It does not "
     "store anything — an item stays in the wall for at most one wave (~3 hours).", sa=4)
para("**Where it is.** In the packing area on the east side of the building, at its south end: right next to the right "
     "perimeter aisle (where the pickers arrive from the fast-pick area and the racks) and just north of the outbound "
     "staging and the shipping doors. Fig 6 shows it on the whole layout and enlarged; it is also drawn on the floor plan (Fig 1).", sa=4)
figure("putwall_where.png", "Fig 6. Where the put-to-light wall is: (a) the whole warehouse, (b) the packing area enlarged, with the key.", 17)
para("**How it works.** The wall is a shelf you can reach from both sides. The pick side faces the aisle, the pack side "
     "faces the pack stations. Fig 7 shows it from the side:", sa=4)
figure("putwall_how.png", "Fig 7. How the wall works, seen from the side: pick side (left), the wall, pack side (right).", 16.5)
bullets([
    f"**Size:** {PWD['columns']} columns × {PWD['levels']} levels = **{PWD['slots']} slots**, {PWD['length_m']} m long, "
    f"{PWD['depth_m']} m deep, slots between 0.2 and 2.0 m high. Each slot holds one open tote (600 × 400 mm, ~40 L) for one order "
    "and has a light with a button on both sides.",
    f"**Why {PWD['slots']} slots:** the peak wave has ~{PWD['small_orders_wave_peak']:.0f} small orders; {TS.get(1, 0):.0f}% need one "
    f"tote, {TS.get(2, 0):.0f}% two and {TS.get(3, 0) + TS.get(4, 0):.0f}% three or four (average {PK['totes_per_small_order']:.2f}). "
    f"{PWD['small_orders_wave_peak']:.0f} × {PK['totes_per_small_order']:.2f} × 1.2 spare → {PWD['slots']} slots. A bigger order gets "
    "2–4 slots next to each other.",
    "**Orders from two zones** wait in their slot until the last item arrives — only then does the pack-side light go on, so "
    "nobody packs half an order. A wrong item gives no light, so picking mistakes are caught here.",
])
para("Where each item is stored before it reaches the wall:", sa=4)
table(["Item type", "Stored in", "How it is stored", "Who picks it"], [
    ["Fast movers (A/B)", "Fast-pick area (south band)", "cases in carton-flow lanes: refilled from the back by the reach truck, "
     "picked from the front at 0.6–1.6 m height", "zone 1 picker with a trolley"],
    ["Slow movers (C)", "Low levels of the reserve racks", "on the pallet, levels 1–2 within reach", "zone 2–3 picker with a trolley"],
    ["Small slow parts", "Bin shelving behind the wall (E)", f"one SKU per half-bin (300 × 400 mm), 7 levels, {SP['shelf_bays']} bays",
     "the putter, between puts"],
    ["Big orders (pallet-sized)", "—", "do not use the wall: picked straight onto a pallet (F)", "EJE 120i picker"],
], widths=[3.0, 3.6, 6.2, 3.2], arr=9)
pk_fte = [r for r in CP["labour_table"] if r["Process"].startswith("Packing")][0]
para("Who works at the wall, and how they move:", sa=4)
table(["Role", "Where (Fig 6)", "How they move", "People"], [
    ["Picker", "zones → aisle H → pick side B", "pushes a trolley along the zone route, leaves the full trolley at B and takes an "
     "empty one back (2 trolleys per picker)", "1–2 per zone per wave"],
    ["Putter", "pick side B + shelving E", f"walks along the {PWD['length_m']} m wall putting items in; picks small parts in E "
     f"(~{PWD['put_s_per_line']:.0f} s per item)", "1 per wave"],
    ["Packer", "pack side C → cages D", "takes a finished order out of the wall, turns to the bench 1–2 m away, puts the carton in "
     "the carrier's cage", f"~{pk_fte['FTE avg']:.1f} FTE (average), ~{pk_fte['FTE peak']:.1f} (peak, 4 stations)"],
], widths=[2.2, 3.6, 7.6, 2.6], arr=9)
doc.add_heading("3.2 How are small items stored efficiently?", level=2)
para(f"Today every SKU sits on a pallet location, even when its whole stock is a few litres. {SP['skus']:,} slow SKUs (C class "
     f"or no demand in Q2) have ≤ 50 litres of stock, units ≤ 20 L, ≤ 15 kg and ≤ 60 cm — together only {SP['stock_m3']:.0f} m³, "
     f"yet they fill {SP['freed_rack_locations']:,} pallet locations of about 1 m³ on their own (and share others). They need {SP['slots_needed']:,} half-bins "
     f"(300 × 400 mm, ~20 L). We compared three ways to store them (A042):", sa=4)
table(["Option", "Floor m²", "Pallet locations used", "One-off cost", "Pick rate", "Cost per year*"],
      [[o["Option"], f"{o['Floor m²']:.0f}", f"{o['Pallet positions used']:,.0f}", eur(o["CAPEX EUR"]), f"{o['Pick rate lines/h']:.0f} lines/h",
        eur(o["Yearly cost EUR"])] for o in PC["small_item_options"]], widths=[4.6, 1.8, 2.4, 2.4, 2.2, 2.6])
para("* depreciation over 7 years + picking labour for the ~" + f"{SP['lines_day']:.0f} small-part lines a day.", size=8.5, color=GREY, italic=True)
bullets([
    f"**Chosen: bin shelving next to the put-wall.** {SP['shelf_bays']} bays, 2.5 m high, in the packing area's spare floor. It is the "
    f"cheapest option, frees {SP['freed_rack_locations']:,} pallet locations, and the picked item goes straight into the wall — no "
    "walk to the reserve aisles.",
    f"**Kardex Shuttle XP 500 (S027):** {SO_LIFT['Notes'].split(';')[0]} would need only {SO_LIFT['Floor m²']:.0f} m² and picks faster with "
    f"a light pointer, but costs ~{eur(SO_LIFT['Yearly cost EUR'] - SO_SHELF['Yearly cost EUR'], True)} a year more. With ~{SP['lines_day']:.0f} "
    "lines a day the lifts would stand still most of the day. They become the better choice when floor space runs out or small-part "
    "volume grows strongly — the shelving bins fit the lift trays, so the step is easy.",
    f"**A/B small items stay in the fast-pick area** ({SP['skus_small_ab_in_fast_pick']} SKUs): they are picked too often for a lift or a "
    "far-away shelf.",
])

# ================================================================ 4 DESIGN CHECKS
doc.add_heading("4. Design checks — machinery, long goods, space and hazardous goods", level=1)
doc.add_heading("4.1 Machinery — is everything fit and justified?", level=2)
para("Each machine was checked on three questions: does it have a job every day, does it fit the loads and the aisles, "
     "and is the number right?", sa=4)
mach = [
    ["ETV 216i reach truck", "Putaway, refill fast-pick area, full pallets", f"3 → {Q['etv']}",
     f"{fleet['ETV 216i']['Utilisation avg %']:.0f}%", "Yes: 10.7 m lift for the 10.0 m top beam; single-deep → standard forks",
     "Keep. 2 cover the peak day; a breakdown still leaves enough for an average day."],
    ["EJE 120i pallet truck", "Large orders picked onto a pallet", f"4 → {Q['eje']}",
     f"{fleet['EJE 120i']['Utilisation avg %']:.0f}%", "Yes: 0.72 m wide, no U-turns needed", "Keep 3 (peak need)."],
    ["ERE 225 rider pallet truck", "Unload / load trailers, container-zone pallets", f"2 → {Q['ere']}",
     f"{fleet['ERE 225']['Utilisation avg %']:.0f}%", "Yes: 2.5 t, carries the few heavy loads",
     "Keep 1 per dock side: docks are 96 m apart and both are busy 10:30–14:00. Low use — could drop to 1 if DHL accepts the walk."],
    ["Pick trolley + totes", "Small orders, batch picking", f"{Q['trolleys']}", "—", "Yes: 250 kg, 6 totes", "Keep (cheap, flexible)."],
    ["Put-to-light wall", "Sort batch picks into orders", f"{Q['ptl_slots']} slots", "—", f"Peak wave ~{SORT['small_orders_wave_peak']:.0f} small orders × {PK['totes_per_small_order']:.2f} totes + 20%", "Keep."],
    ["Small-parts bin shelving", "Store and pick small slow parts", f"{SP['shelf_bays']} bays", "—", "2.5 m high, 7 levels; within reach without a ladder", "Add (feedback 2)."],
    ["Roller conveyor + 2 stretch wrappers", "Unpack containers; wrap new and outgoing pallets", "new", "—",
     "Yes: no heavy carrying out of the container", "Add."],
    ["RF terminals", "Scan every move", f"{Q['rf_units']}", "—", "One per person at the peak + 20%", "Keep."],
    ["Telescopic forks", "(double-deep)", "dropped", "—", "Not needed with single-deep racks", "Removed."],
    ["Cantilever racking", "(items > 2.4 m)", "dropped", "—", "No real item is longer than ~2 m", "Removed."],
    ["Kardex lifts", "(small parts)", "compared", "—", "Fit, but little work at our volume", "Upgrade path (§3.2)."],
    ["ECE 225 order picker / EZS 130 tugger", "(rejected in D2)", "—", "—", "Too long for 3.0 m aisles / driver idle", "Not used."],
]
table(["Machine", "Job", "Units", "Use (avg day)", "Fits?", "Verdict"], mach, widths=[2.8, 2.8, 1.6, 1.4, 3.6, 3.8], arr=9)

doc.add_heading("4.2 Long goods — lengths, turning, capacity and access", level=2)
bullets([
    f"**Real lengths:** nothing longer than a door (~2 m) can come in. The WMS shows {LG['loads_over_2_05m_data_error']} loads longer "
    "than 2.05 m — up to an '11 m' item that weighs 0.5 kg and is 1 cm high — so these are data errors and are capped at 2.05 m, "
    f"the length of an XLONG pallet. That gives {LG['loads_long_capped']} long loads (1.3–2.05 m).",
    f"**Where they go:** the long-goods module has 4.5 m bays with two long pallets each: {LAY['positions_long']} positions for "
    f"{LG['loads_long_capped']} loads today ({LG['loads_long_capped'] / LAY['positions_long']:.0%} full).",
    f"**Turning clearance:** a reach truck stacking a 2.05 m pallet crosswise needs an aisle of {LG['ast_long_crosswise_mm']:,.0f} mm "
    f"(VDI right-angle stacking formula with the ETV 216i data: Wa 1,665 mm, 200 mm safety, S006). A 3.0 m aisle leaves only "
    f"{3000 - LG['ast_long_crosswise_mm']:.0f} mm, so the long-goods aisle is 3.2 m ({3200 - LG['ast_long_crosswise_mm']:.0f} mm spare). "
    "The 3.2 m cross-aisles let the truck turn into it with the load.",
    f"**Truck capacity:** the forks go under the long side, so the load centre is ~{LG['load_centre_long_mm']} mm — inside the "
    f"600 mm rating, full 1.6 t. Forked on the short side the load centre would be ~1 m and the truck could carry only "
    f"~{LG['capacity_if_forked_short_side_kg']:.0f} kg; the heaviest long load is {LG['long_load_kg_max']:,.0f} kg (95% below "
    f"{LG['long_load_kg_p95']:.0f} kg), so loads over 1.6 t go to a floor position.",
    "**Access:** the module is open from both ends (front, middle and rear cross-aisles) and its outer row faces the 3.2 m "
    "perimeter aisle, so long goods never pass through the standard 3.0 m aisles.",
])
para("Loads checked against the racks and trucks (all 11,015 loads in stock, 01-06-2025):", sa=4)
table(["Check", "Data", "Result"], [
    ["Weight", f"{LC['weight']['over_1000']} loads > 1 t, {LC['weight']['over_1600']} > 1.6 t (actual qty × unit weight)",
     "ETV 216i lifts 1.6 t → the few heavier loads go to floor positions, moved with the ERE 225 (2.5 t)"],
    ["Height", f"{LC['tall']['loads_1_16_to_1_68']} loads 1.16–1.68 m tall, {LC['tall']['loads_over_1_68']} taller; "
               f"{LC['tall']['height_missing']} have no height in the WMS",
     f"tall loads use the top level ({LC['tall']['top_level_positions']:,} positions with ~1.7 m free); the tallest stand on the floor"],
    ["Aisles", "3.0 m storage, 3.2 m long-goods / cross / perimeter", "ETV needs 2.63–2.79 m (2.98 m with a long load), EJE 1.79–1.94 m → OK"],
], widths=[2.0, 6.2, 7.8], arr=9)

doc.add_heading("4.3 Space for unpacking loose containers", level=2)
bullets([
    f"**What a container needs:** cartons are sorted on the floor into one build spot per SKU. Measured on {LZ['containers_sampled']} "
    f"Asian receipts: the median container needs ~{LZ['need_m2_median']:.0f} m², the p90 ~{LZ['need_m2_p90']:.0f} m², the largest "
    f"({LZ['skus_per_container_max']} SKUs) ~{LZ['need_m2_max']:.0f} m².",
    f"**The zone ({LZ['area_planned_m2']} m²) fits {LZ['share_fit_zone_pct']:.0f}% of containers.**",
    f"**Borrowing door R2 (feedback idea, A044):** receiving needs at most ~{LZ['recv_need_m2']:.0f} m² of its 262 m², because two "
    f"reach trucks put pallets away while a trailer is unloaded (peak ~{LZ['recv_staged_pallets_peak']:.0f} pallets waiting). On "
    f"a big-container day R2 is blocked and {LZ['r2_lend_m2']:.0f} m² of receiving joins the zone → {LZ['area_with_r2_m2']:.0f} m², "
    f"enough for {LZ['share_fit_with_r2_pct']:.0f}% of containers. R1 + R3 still give the {LZ['doors_needed_base']} doors the base "
    "and peak days need; R2 is only blocked when no second trailer is booked.",
    "**The last few %** (more SKUs than build spots) are sorted in two passes.",
])
doc.add_heading("4.4 Outbound staging (216 m²) and packing (450 m²)", level=2)
bullets([
    f"**Outbound staging:** all afternoon trailers' goods wait at the same time after wave 3. Measured per day: median "
    f"{OA['need_m2_median']:.0f} m², p90 {OA['need_m2_p90']:.0f} m², busiest day {OA['need_m2_max']:.0f} m² → every Q2 day fits, "
    "with room for twice the volume.",
    f"**Packing:** pack stations, put-wall, small-parts shelving, waiting pallets, wrapper and packaging stock need "
    f"{PA['need_m2']:.0f} m² → fits in 450 m² with {PA['spare_m2']:.0f} m² spare (D2 §5.4).",
])
doc.add_heading("4.5 Hazardous goods — recount and classes", level=2)
para(f"Counted per load (LODNUM) with the hazmat flag: {HZ['skus']} SKUs and {HZ['loads']:,} loads ({HZ['share_loads_pct']:.0f}% of "
     f"loads), ~{HZ['kg_in_stock'] / 1000:.0f} t gross on 01-06-2025 (same method in April and May: "
     f"{', '.join(str(x['loads']) for x in HZ['by_snapshot'][:2])} loads — the count grows with the stock). The customs code "
     "tells what the products are:", sa=4)
table(["Customs code", "What it is", "Likely hazard", "Dutch rule", "Loads", "Tonnes"],
      [[c["HS"], c["What it is"], c["Likely hazard"], c["Dutch rule"], f"{c['Loads']:,}", f"{c['t']:.0f}"] for c in HZ["classes"]],
      widths=[1.6, 3.8, 4.6, 2.0, 1.4, 1.4], arr=4)
bullets([
    f"**{HZ['pgs37_share_loads_pct']:.0f}% of hazardous loads are electronic security products** (safes, locks, alarms, cameras, readers) "
    "— most likely with lithium cells inside (UN3481/UN3091, class 9). Storage then follows PGS 37-2 (from 333 kg; 1,000 kg for "
    "small cells, S029), not PGS 15.",
    "**Only a few loads are chemicals** (lubricants → PGS 15 if aerosols or flammable), and 33 loads of cardboard boxes look like a "
    "wrong flag.",
    "**To confirm with Assa Abloy:** the UN numbers on the safety data sheets. The hazmat section and its fire separation are sized "
    "after that and are not in this budget.",
])

# ================================================================ 5 COSTS
doc.add_heading("5. Costs — CAPEX and OPEX", level=1)
para("The capacity plan (D2 §6) is turned into money in three steps: (1) count what the design needs (positions, bays, "
     "trucks, people), (2) multiply by a price from a public source or a labelled assumption (Section 7), (3) add it up per "
     f"year. Labour is costed at €{RATE['value']:.0f} per hour for an operator (feedback; €{RATE['high']:.0f} tested as the high case). "
     "Trucks are leased with full service (repairs included), so they are OPEX. The building, dock levellers, sprinklers and "
     "the hazmat compartment are outside the project scope and not included.", sa=6)
doc.add_heading("5.1 One-off investment (CAPEX)", level=2)
figure("d3_capex.png", f"Fig 8. CAPEX by group — total {eur(PC['capex_total'], True)}.", 13.5)
rows = [[r["Item"], f"{r['Qty']:,.0f} {r['Unit']}", eur(r["Unit price EUR"]) if r["Unit price EUR"] >= 10 else f"€{r['Unit price EUR']:.2f}",
         eur(r["Total EUR"]), r["Source"]] for r in PC["capex_table"]]
rows.append(["Total CAPEX", "", "", eur(PC["capex_total"]), ""])
table(["Item", "Quantity", "Unit price", "Total", "Source"], rows, widths=[6.0, 2.8, 2.0, 2.4, 2.8], arr=2, bold_last=True)
bullets([
    f"**Racking is {PC['capex_by_group']['Storage'] / PC['capex_total']:.0%} of the investment.** If the racking already standing in "
    f"Bemmel can be rebuilt in the new layout (about €{PC['cost_register']['rack_reuse']['value']:.0f} per position), the investment drops to "
    f"{eur(PC['capex_reuse_racking'], True)}. DHL must confirm the type and condition of today's racks.",
    f"**Depreciation:** {eur(PC['depreciation_year'])} a year (racking 15 years, equipment 7, IT 5).",
])
doc.add_heading("5.2 Running cost (OPEX)", level=2)
figure("d3_opex.png", f"Fig 9. OPEX by group — total {eur(PC['opex_total'], True)} per year.", 13.5)
rows = [[r["Item"], f"{r['Qty / yr']:,.0f} {r['Unit']}", f"€{r['Unit cost EUR']:,.2f}" if r["Unit cost EUR"] < 10 else eur(r["Unit cost EUR"]),
         eur(r["Total EUR/yr"]), r["Source"]] for r in PC["opex_table"]]
rows.append(["Total OPEX per year", "", "", eur(PC["opex_total"]), ""])
rows.append(["+ depreciation of CAPEX", "", "", eur(PC["depreciation_year"]), "A041"])
rows.append(["Total yearly cost", "", "", eur(PC["annual_cost"]), ""])
table(["Item", "Quantity / year", "Unit cost", "Per year", "Source"], rows, widths=[6.0, 2.8, 2.0, 2.4, 2.8], arr=2, bold_last=True)
para(f"Labour hours come from the capacity plan: {STO['fte_avg']:.1f} FTE on an average day × 8 paid hours × "
     f"{PC['cost_register']['days_year']['value']} working days. Peak days are covered by flexible staff at the same hourly cost.",
     size=9.5, color=GREY)
doc.add_heading("5.3 Cost per unit", level=2)
uc = PC["unit_cost"]; vy = PC["volumes_year"]
table(["Per …", "Volume per year", "Cost"], [
    ["Order line", f"{vy['lines']:,.0f}", f"€{uc['per_order_line']:.2f}"],
    ["Order", f"{vy['orders']:,.0f}", f"€{uc['per_order']:.2f}"],
    ["Inbound pallet", f"{vy['inbound_pallets']:,.0f}", f"€{uc['per_inbound_pallet']:.2f}"],
    ["Pallet position per month", f"{LAY['positions_total']:,}", f"€{uc['per_position_month']:.2f}"],
], widths=[5.0, 4.0, 3.0])
para("These are full costs (OPEX + depreciation) split over the volume — useful for DHL to compare with the price "
     "agreed with Assa Abloy.", size=9.5, color=GREY)
doc.add_heading("5.4 Today vs future — saving and payback", level=2)
table(["", "Process changes only", "Whole design, re-used racking", "Whole design, all new"], [
    ["One-off investment", eur(PC["capex_process"]), eur(PC["capex_reuse_racking"]), eur(PC["capex_total"])],
    ["Labour saved per year (Section 2.2)", eur(PC["saving_year"]), eur(PC["saving_year"]), eur(PC["saving_year"])],
    ["Simple payback", f"{PC['payback_process_years']:.1f} years", f"{PC['payback_reuse_years']:.1f} years", f"{PC['payback_years']:.1f} years"],
], widths=[5.0, 3.6, 3.8, 3.6])
bullets([
    "**Process changes** = carton-flow fast-pick area, trolleys and totes, put-to-light wall, RF terminals and the WMS set-up of "
    "waves and slotting. They create the whole labour saving and pay for themselves quickly — they can start before the racks are rebuilt.",
    "**Racks, shelving and the container zone are capacity**: they are needed to store and receive the goods at all, so they are "
    "judged on cost per position and space, not on labour saved.",
    "**Only labour savings are counted.** Shorter lead times, fewer parked trailers, fewer picking errors and a safer layout have "
    "value too, but no reliable price.",
])
doc.add_heading("5.5 What if our cost assumptions are wrong?", level=2)
figure("d3_tornado.png", "Fig 10. Yearly cost (OPEX + depreciation) when each assumption moves to the low or high end of its range.", 15)
bullets([
    f"**Labour rate and productivity matter most** — labour is ~{PC['unit_cost']['labour_share_pct']:.0f}% of the yearly cost. DHL's real "
    "hourly cost and time studies (A006, A025) are the most valuable inputs to get.",
    "**Racking price is the biggest one-off uncertainty**, but spread over 15 years its effect on the yearly cost is smaller.",
    f"**The truck lease price hardly matters**: only {Q['etv'] + Q['eje'] + Q['ere']} trucks are leased.",
    f"**Loose containers:** a DHL crew costs ~€{PC['loose_cost_per_container']:,.0f} per container "
    f"({Q['containers_year']:.0f} containers a year). If a contractor quotes less, use the contractor in the same zone.",
])

# ================================================================ 6 SUPPORTING DOCUMENTS
doc.add_heading("6. Supporting documents", level=1)
table(["Document", "What it contains"], [
    ["Deliverable 1 — Analysis (updated)", "Data analysis, concept choice and decision matrix; now single-deep racks and per-order lead time"],
    ["Deliverable 2 — Future state (v3)", "Processes, flowcharts, layout and 3D, docks, picking, machinery, capacity plan; now with the container "
     "zone and R2 overflow, long-goods module, small-parts shelving, area checks and the hazmat classes"],
    ["Deliverable_3_Calculations.xlsx", "Every price, quantity and formula of this report; change an input and all totals update"],
    ["data/assumptions.json", "All assumptions A001–A045, each with reason, source and impact"],
    ["documentation/sources.md", "All sources S001–S030 with links and what they support"],
    ["analysis/*.py", "Scripts that turn the raw data into every table and figure (run order in README.md)"],
], widths=[5.4, 10.6], arr=9)

# ================================================================ 7 ASSUMPTIONS
doc.add_heading("7. Cost assumptions and open points", level=1)
para("Assumptions added for Deliverable 3 (full register in data/assumptions.json):", sa=4)
rows = [[a["id"], a["assumption"], a["value"], a["source"]] for a in ASSUM["assumptions"] if a["id"] >= "A030"]
table(["ID", "Assumption", "Value", "Source"], rows, widths=[1.2, 3.6, 7.4, 3.8], arr=9)
para("Sources for prices and benchmarks:", sa=4)
src = {
    "S018": "CBS — labour cost per worked hour 2024 (transport & storage)",
    "S019": "Cao Beroepsgoederenvervoer — wage scales 2025",
    "S020": "IndexBox / T-Racking — pallet racking price per position",
    "S021": "Dutch reach-truck lease and rental offers (Liftlease, Beekmans, Itolang, Solid Lift)",
    "S022": "Eurostat nrg_pc_205 — NL non-household electricity price H1 2025",
    "S023": "WERC DC Measures 2024/2025 — pick accuracy and dock-to-stock benchmarks",
    "S024": "Prestatie — loose-load container unloading (2–4 h, fixed price per container)",
    "S025": "Public equipment listings (trolley, tote, wrapper, conveyor, carton flow, put-to-light, RF terminal)",
    "S026": "Teacher feedback session 2 — space checks, ~2 m max length, small parts vs Kardex, own KPIs, €20/h, hazmat",
    "S027": "Relevator Group — Kardex Shuttle XP 500 specifications and prices",
    "S028": "Cisco-Eagle / Kardex LogiMat / MMH Van Meter — picking rates at lifts and shelving",
    "S029": "Evofenedex / Specialcargo / BOVAG — PGS 37-2 storage of lithium batteries",
    "S030": "Manutan — steel shelving section prices",
}
table(["ID", "Source"], [[k, v] for k, v in src.items()], widths=[1.4, 14.6], arr=9)
bullets([
    "**To get from DHL:** real hourly labour cost and pick/pack times (A006, A025, A030); type and condition of today's racking "
    "(re-use); container numbers per week (A019).",
    "**To get from suppliers:** Jungheinrich lease quote for the trucks (A038); racking quote incl. the long-goods module (A033); "
    "contractor price per loose container (A026).",
    "**To get from Assa Abloy:** UN numbers / safety data sheets of the hazardous SKUs (A028), and a clean-up of the item lengths "
    "in the master data (A043).",
])

footer()
DELIV = os.path.join(BASE, "deliverables"); os.makedirs(DELIV, exist_ok=True)
out = os.path.join(DELIV, "Deliverable_3_Performance_Overview.docx")
doc.save(out)
print("Saved:", out)
