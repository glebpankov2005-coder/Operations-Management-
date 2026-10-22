"""
Who is responsible for what (teacher: every student hands in the whole design and marks the process they own).
One place for the owner of each process and of each report section, used by all report builders and the portfolio.
Change a name or a section here and every document follows.
"""
from docx.shared import Pt, RGBColor

TEAM = {"Receiving": "Vidhi", "Storage": "Gleb", "Picking": "Arnau", "Shipping": "Rachel"}
PROCESS_SCOPE = {
    "Receiving": "unloading trailers and containers, checking, container unpacking, returns intake — up to the pick-up point",
    "Storage": "putaway, racks and storage places, slotting, replenishment of the fast-pick area, inventory control",
    "Picking": "order release (waves), picking, sorting batch picks into orders at the put-to-light wall",
    "Shipping": "packing, outbound staging, dock doors, trailer departure plan and loading",
}
COLOR = {"Receiving": RGBColor(0xC2, 0x41, 0x0C), "Storage": RGBColor(0x1D, 0x4E, 0xD8),
         "Picking": RGBColor(0x15, 0x80, 0x3D), "Shipping": RGBColor(0x7C, 0x3A, 0xED)}
HEX = {"Receiving": "C2410C", "Storage": "1D4ED8", "Picking": "15803D", "Shipping": "7C3AED"}

# heading number -> processes it belongs to (sections not listed were written by the whole group)
OWNERS = {
    "D1": {"3.": ["Storage"], "4.": ["Storage"], "5.": ["Receiving"], "6.": ["Picking"], "7.": ["Shipping"],
           "8.": ["Storage", "Picking"], "9.": ["Picking", "Shipping"], "13.": ["Storage"]},
    "D2": {"2.1": ["Receiving", "Storage"], "2.2": ["Storage"], "2.3": ["Picking"], "2.4": ["Shipping", "Picking"],
           "2.5": ["Receiving"], "3.1": ["Storage"], "3.2": ["Storage"], "4.1": ["Receiving"], "4.2": ["Picking"],
           "4.3": ["Receiving"], "4.4": ["Shipping"], "4.5": ["Shipping"], "4.6": ["Receiving"], "5.2": ["Storage"],
           "5.3": ["Storage"], "5.4": ["Shipping"], "6.2": ["Picking"], "7.1": ["Receiving"], "7.2": ["Storage"],
           "7.3": ["Picking"], "7.4": ["Shipping"]},
    "D3": {"2.1": ["Shipping"], "2.2": ["Picking"], "2.3": ["Storage"], "3.1": ["Receiving"], "4.": ["Storage"],
           "4.1": ["Storage"], "4.2": ["Storage"], "4.3": ["Storage", "Picking"], "5.1": ["Picking", "Shipping"],
           "5.2": ["Storage"], "6.2": ["Storage"], "6.3": ["Receiving"], "6.4": ["Shipping"], "6.5": ["Storage"]},
}
DOC_NAME = {"D1": "Deliverable 1 — Analysis", "D2": "Deliverable 2 — Future state", "D3": "Deliverable 3 — Performance overview"}


def tag(processes):
    return " · ".join(f"{TEAM[p]} ({p})" for p in processes)


def _number(text):
    first = text.split(" ", 1)[0]
    return first if first[:1].isdigit() else None


def tag_headings(doc, which):
    """Add a coloured owner tag to every heading that belongs to one process. Returns {process: [heading texts]}."""
    owned = {p: [] for p in TEAM}
    for par in doc.paragraphs:
        if not par.style.name.startswith("Heading") or par.style.name == "Heading 1" and par.text.strip() == "Contents":
            continue
        num = _number(par.text)
        procs = OWNERS[which].get(num)
        if not procs:
            continue
        for i, p in enumerate(procs):
            r = par.add_run(("   " if i == 0 else "  ·  ") + f"{TEAM[p]} ({p})")
            r.font.size = Pt(9.5); r.font.color.rgb = COLOR[p]; r.bold = True
        for p in procs:
            owned[p].append(par.text.split("   ")[0])
    return owned


def sections_of(which, process):
    return [k.rstrip(".") for k, v in OWNERS[which].items() if process in v]


def responsibility_rows(which):
    """Rows for the 'who is responsible for what' table at the start of a deliverable."""
    return [[p, TEAM[p], PROCESS_SCOPE[p], ", ".join("§" + s for s in sections_of(which, p))] for p in TEAM]


def responsibility_box(doc, which, intro=True):
    """'Who is responsible for what' block (cover page of each deliverable)."""
    from docx.shared import Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = doc.add_paragraph(); r = p.add_run("Who is responsible for what"); r.bold = True; r.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    if intro:
        p = doc.add_paragraph()
        r = p.add_run("Every team member hands in this whole document (the complete warehouse design). Each process has one "
                      "owner; the sections of a process carry the owner's name in their heading. Sections without a name "
                      "were made by the whole group.")
        r.font.size = Pt(9); r.italic = True; p.paragraph_format.keep_with_next = True
    rows = responsibility_rows(which)
    t = doc.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Process", "Owner", "Scope", f"Sections in this document"]):
        t.rows[0].cells[i].text = ""; rr = t.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
    for proc, name, scope, secs in rows:
        cells = t.add_row().cells
        for i, v in enumerate([proc, name, scope, secs]):
            cells[i].text = ""; rr = cells[i].paragraphs[0].add_run(v); rr.font.size = Pt(9)
            if i < 2:
                rr.bold = True; rr.font.color.rgb = COLOR[proc]
    widths = [2.4, 1.8, 7.8, 4.0]
    grid = t._tbl.tblGrid
    for i, w in enumerate(widths):
        grid.gridCol_lst[i].w = Cm(w)
        for row in t.rows:
            row.cells[i].width = Cm(w)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
    return t


def update_fields_on_open(doc):
    """Ask Word to refresh the table of contents and page fields when the file is opened."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")):
        settings.remove(el)
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    settings.append(uf)


FULL_NAMES = {"Vidhi": "Vidhi", "Gleb": "Gleb", "Arnau": "Arnau", "Rachel": "Rachel Selvarathnam"}


def team_line():
    return "Group 3 — " + ", ".join(f"{FULL_NAMES[n]} ({p})" for p, n in TEAM.items())
