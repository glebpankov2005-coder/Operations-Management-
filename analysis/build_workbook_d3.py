"""
Deliverable 3 calculation workbook (Excel, live formulas).
Every cost in the D3 report is rebuilt here from named inputs, so a reader can change a price or a rate on the
'Inputs' sheet and see CAPEX, OPEX, unit costs and payback update. Quantities and workloads come from the
analysis (phase1_metrics.json); prices from the cost register in performance_cost.py (A030-A041).
Output: deliverables/Deliverable_3_Calculations.xlsx. Run after performance_cost.py.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), "..")
M = json.load(open(os.path.join(BASE, "outputs", "reports", "phase1_metrics.json"), encoding="utf-8"))
PC, DM, CP, LAY = M["performance_cost"], M["dock_mhe"], M["capacity_plan"], M["layout"]
REG = PC["cost_register"]; Q = PC["quantities"]

wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F3A68"); INPUT = PatternFill("solid", fgColor="FFF4CC"); TOT = PatternFill("solid", fgColor="E3EBF8")
WHITE = Font(color="FFFFFF", bold=True); BOLD = Font(bold=True); THIN = Side(style="thin", color="C7CEDE")
EUR = '€#,##0;-€#,##0'; EUR2 = '€#,##0.00'; NUM = '#,##0'; NUM1 = '#,##0.0'; PCT = '0%'


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c); cell.fill = HEAD; cell.font = WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def name(ws, ref, nm):
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws.title}'!{ref}")


# ------------------------------------------------------------------ Read me
ws = wb.active; ws.title = "Read me"
lines = [
    ("Project OTM — Assa Abloy @ DHL Bemmel · Deliverable 3 calculations", True),
    ("", False),
    ("How to use: change a yellow cell on 'Inputs' (prices, rates) or 'Volumes' (workload). CAPEX, OPEX and Summary update.", False),
    ("Yellow = input. Blue rows = totals. All other numbers are formulas or values copied from the analysis scripts.", False),
    ("Volumes are measured in the Assa Abloy Q2-2025 data (analysis/*.py). Prices are public benchmarks or labelled", False),
    ("assumptions (data/assumptions.json A030-A041, sources S018-S025 in documentation/sources.md) — not quotes.", False),
    ("", False),
    ("Sheets: Summary · Inputs · Volumes · CAPEX · OPEX · Labour · KPIs · Small parts · Stress test · Sensitivity", False),
]
for r, (t, b) in enumerate(lines, 1):
    ws.cell(row=r, column=1, value=t).font = Font(bold=b, size=13 if b else 11)
ws.column_dimensions["A"].width = 120

# ------------------------------------------------------------------ Inputs
wi = wb.create_sheet("Inputs")
header(wi, 1, ["Name", "Value", "Low", "High", "Unit", "ID", "Basis / source"])
for r, (k, v) in enumerate(REG.items(), 2):
    wi.cell(row=r, column=1, value=k)
    c = wi.cell(row=r, column=2, value=v["value"]); c.fill = INPUT
    wi.cell(row=r, column=3, value=v["low"]); wi.cell(row=r, column=4, value=v["high"])
    wi.cell(row=r, column=5, value=v["unit"]); wi.cell(row=r, column=6, value=v["id"]); wi.cell(row=r, column=7, value=v["basis"])
    name(wi, f"$B${r}", k)
widths(wi, [16, 11, 9, 9, 42, 7, 120])
wi.freeze_panes = "A2"

# ------------------------------------------------------------------ Volumes (workload from the analysis)
wv = wb.create_sheet("Volumes")
labour = [r for r in CP["labour_table"] if r["DHL staff"]]
fleet = {f["Model"].replace("Jungheinrich ", ""): f for f in DM["fleet"]}
t_put = {c["Move"]: c for c in DM["cycle_times"]}["Putaway (receiving -> reserve)"]["Cycle min"]
vol = [
    ("dir_h_day", sum(r["Hours (avg)"] for r in labour), "productive DHL labour hours per average day", "capacity plan (D2 §6.1)"),
    ("loose_crew", DM["loose_zone"]["crew"], "people in the loose-container crew", "A019"),
    ("loose_h_cont", DM["loose_zone"]["hours_per_container"], "hours per container", "A019"),
    ("cont_day", CP["inbound_pattern"]["containers_avg"], "loose containers per working day", "A019"),
    ("lines_day", M["peak"]["lines"]["mean"], "order lines per working day", "measured"),
    ("orders_day", M["peak"]["orders"]["mean"], "orders per working day", "measured"),
    ("inb_day", CP["measured"]["inbound_loads_day"], "inbound pallets per working day", "measured"),
    ("positions", LAY["positions_total"], "pallet positions (single-deep, incl. long goods)", "floor plan"),
    ("positions_long", LAY["positions_long"], "long-goods positions (4.5 m bays)", "floor plan"),
    ("n_etv", Q["etv"], "ETV 216i reach trucks", "fleet review (D3 §4)"),
    ("n_eje", Q["eje"], "EJE 120i pallet trucks", "fleet review"),
    ("n_ere", Q["ere"], "ERE 225 rider pallet trucks", "fleet review"),
    ("etv_h_day", fleet["ETV 216i"]["Equip-h/day avg"], "ETV operating hours per day", "cycle-time model"),
    ("pt_h_day", fleet["EJE 120i"]["Equip-h/day avg"] + fleet["ERE 225"]["Equip-h/day avg"], "pallet-truck operating hours per day", "cycle-time model"),
    ("t_put_min", t_put, "putaway cycle, minutes per pallet", "cycle-time model"),
    ("loads_now", max(s["loads_LODNUM"] for s in M["inventory"]["snapshots"]), "loads to move in once (re-slotting)", "measured"),
    ("pick_today_h", PC["pick_today_staffed_h"], "picker-hours/day today (modelled, incl. 08:00 peak staffing)", "D3 §2.2"),
    ("pick_future_h", PC["pick_future_h"], "picker-hours/day in the design", "pick-time model"),
]
header(wv, 1, ["Name", "Value", "What", "From"])
for r, (k, v, what, src) in enumerate(vol, 2):
    wv.cell(row=r, column=1, value=k); c = wv.cell(row=r, column=2, value=v); c.fill = INPUT; c.number_format = NUM1
    wv.cell(row=r, column=3, value=what); wv.cell(row=r, column=4, value=src)
    name(wv, f"$B${r}", k)
widths(wv, [16, 12, 60, 28])

# ------------------------------------------------------------------ CAPEX
wc = wb.create_sheet("CAPEX")
header(wc, 1, ["Group", "Item", "Qty", "Unit", "Unit price", "Total", "Life (yrs)", "Depreciation / yr", "Source"])
price_formula = {   # item prefix -> unit-price formula (named inputs)
    "Single-deep pallet racking": "=rack_pp", "Long-goods racking": "=long_bay_pp", "Rack protection": "=rack_protect*positions*rack_pp",
    "Bin shelving": "=shelf_bay", "Half-bins": "=vlm_slot",
    "Carton-flow": "=flow_bay", "Pick trolleys": "=trolley", "Totes": "=tote", "Put-to-light wall": "=ptl_slot",
    "Put-to-light installation": "=ptl_integr", "Pack stations": "=pack_station", "Stretch wrapper": "=wrapper",
    "Roll cages": "=roll_cage", "Extendable roller": "=conveyor", "RF terminals": "=rf_unit", "WMS set-up": "=wms_config",
    "Move-in": "=t_put_min/60*(labour_rate+lease_etv*12/(days_year*7.5))",
}
life_name = {"Storage": "life_rack", "IT": "life_it"}
rows = [r for r in PC["capex_table"] if r["Group"] != "Contingency"]
r0 = 2
for i, it in enumerate(rows):
    r = r0 + i
    wc.cell(row=r, column=1, value=it["Group"]); wc.cell(row=r, column=2, value=it["Item"])
    q = wc.cell(row=r, column=3, value="=positions-positions_long" if it["Item"].startswith("Single-deep")
                else ("=positions_long" if it["Item"].startswith("Long-goods") else it["Qty"])); q.number_format = NUM
    wc.cell(row=r, column=4, value=it["Unit"])
    f = next((v for k, v in price_formula.items() if it["Item"].startswith(k)), it["Unit price EUR"])
    wc.cell(row=r, column=5, value=f).number_format = EUR2
    wc.cell(row=r, column=6, value=f"=C{r}*E{r}").number_format = EUR
    wc.cell(row=r, column=7, value=f"={life_name.get(it['Group'], 'life_equip')}")
    wc.cell(row=r, column=8, value=f"=F{r}/G{r}").number_format = EUR
    wc.cell(row=r, column=9, value=it["Source"])
rl = r0 + len(rows) - 1
rc = rl + 1
wc.cell(row=rc, column=1, value="Contingency"); wc.cell(row=rc, column=2, value="Contingency on all items above")
wc.cell(row=rc, column=3, value=1); wc.cell(row=rc, column=4, value="lot")
wc.cell(row=rc, column=5, value=f"=contingency*SUM(F{r0}:F{rl})").number_format = EUR
wc.cell(row=rc, column=6, value=f"=C{rc}*E{rc}").number_format = EUR
wc.cell(row=rc, column=7, value="=life_equip"); wc.cell(row=rc, column=8, value=f"=F{rc}/G{rc}").number_format = EUR
wc.cell(row=rc, column=9, value="A037")
rt = rc + 1
wc.cell(row=rt, column=2, value="TOTAL CAPEX").font = BOLD
wc.cell(row=rt, column=6, value=f"=SUM(F{r0}:F{rc})").number_format = EUR
wc.cell(row=rt, column=8, value=f"=SUM(H{r0}:H{rc})").number_format = EUR
for c in range(1, 10): wc.cell(row=rt, column=c).fill = TOT
name(wc, f"$F${rt}", "capex_total"); name(wc, f"$H${rt}", "depreciation")
rack_rows = [r0 + i for i, it in enumerate(rows) if it["Item"].startswith(("Single-deep", "Long-goods", "Rack protection"))]
proc_rows = [r0 + i for i, it in enumerate(rows) if it["Item"].startswith(("Carton-flow", "Pick trolleys", "Totes", "Put-to-light", "WMS set-up", "RF terminals"))]
rr = rt + 2
wc.cell(row=rr, column=2, value="CAPEX if today's racking is re-used (dismantle + rebuild at rack_reuse per position)").font = BOLD
wc.cell(row=rr, column=6, value=f"=(SUM(F{r0}:F{rl})-{'-'.join('F' + str(x) for x in rack_rows)}+positions*rack_reuse)*(1+contingency)").number_format = EUR
for c in range(1, 10): wc.cell(row=rr, column=c).fill = TOT
name(wc, f"$F${rr}", "capex_reuse")
rp = rr + 1
wc.cell(row=rp, column=2, value="CAPEX of the process changes only (fast-pick, trolleys, totes, put-wall, RF, WMS set-up)").font = BOLD
wc.cell(row=rp, column=6, value=f"=({'+'.join('F' + str(x) for x in proc_rows)})*(1+contingency)").number_format = EUR
for c in range(1, 10): wc.cell(row=rp, column=c).fill = TOT
name(wc, f"$F${rp}", "capex_process")
widths(wc, [18, 52, 9, 12, 13, 14, 9, 15, 22])
wc.freeze_panes = "A2"

# ------------------------------------------------------------------ OPEX
wo = wb.create_sheet("OPEX")
header(wo, 1, ["Group", "Item", "Qty / yr", "Unit", "Unit cost", "Total / yr", "Formula in words"])
opex_rows = [
    ("Labour", "Direct DHL labour", "=dir_h_day/7.5*paid_h_shift*days_year", "paid h", "=labour_rate", "productive h ÷ 7.5 × 8 paid h × days"),
    ("Labour", "Loose-container zone crew", "=loose_crew*loose_h_cont*cont_day/7.5*paid_h_shift*days_year", "paid h", "=labour_rate", "crew × h per container × containers/day"),
    ("Labour", "Team leader + inventory control / admin", "=indirect_fte*paid_h_shift*days_year", "paid h", "=(lead_rate+0.5*admin_rate)/1.5", "1 lead + 0.5 admin"),
    ("Machinery", "ETV 216i lease", "=n_etv*12", "truck-months", "=lease_etv", "full-service lease"),
    ("Machinery", "EJE 120i lease", "=n_eje*12", "truck-months", "=lease_eje", "full-service lease"),
    ("Machinery", "ERE 225 lease", "=n_ere*12", "truck-months", "=lease_ere", "full-service lease"),
    ("Machinery", "Electricity for trucks", "=(etv_h_day*kwh_truck_h+pt_h_day*1.5)*days_year", "kWh", "=power_price", "operating h × kWh/h × price"),
    ("Upkeep", "Rack inspection and repairs", 1, "per year", "=rack_inspect", "EN 15635"),
    ("Upkeep", "IT support, licences, spares", 1, "per year",
     "=it_support*(SUMIF(CAPEX!A:A,\"IT\",CAPEX!F:F)+SUMIF(CAPEX!B:B,\"Put-to-light*\",CAPEX!F:F))", "share of IT + put-to-light CAPEX"),
]
for i, (g, it, q, u, p, txt) in enumerate(opex_rows, 2):
    wo.cell(row=i, column=1, value=g); wo.cell(row=i, column=2, value=it)
    wo.cell(row=i, column=3, value=q).number_format = NUM; wo.cell(row=i, column=4, value=u)
    wo.cell(row=i, column=5, value=p).number_format = EUR2
    wo.cell(row=i, column=6, value=f"=C{i}*E{i}").number_format = EUR; wo.cell(row=i, column=7, value=txt)
ot = len(opex_rows) + 2
wo.cell(row=ot, column=2, value="TOTAL OPEX per year").font = BOLD
wo.cell(row=ot, column=6, value=f"=SUM(F2:F{ot - 1})").number_format = EUR
for c in range(1, 8): wo.cell(row=ot, column=c).fill = TOT
name(wo, f"$F${ot}", "opex_total"); name(wo, "$F$3", "loose_cost_year")
name(wo, "$F$2:$F$4", "labour_rows")
widths(wo, [12, 40, 12, 13, 12, 14, 42])

# ------------------------------------------------------------------ Summary
wsu = wb.create_sheet("Summary", 1)
header(wsu, 1, ["Result", "Value", "Formula in words"])
summ = [
    ("One-off investment (CAPEX), all new", "=capex_total", EUR, "CAPEX sheet"),
    ("One-off investment if today's racking is re-used", "=capex_reuse", EUR, "CAPEX sheet"),
    ("Running cost (OPEX) per year", "=opex_total", EUR, "OPEX sheet"),
    ("Depreciation per year", "=depreciation", EUR, "CAPEX ÷ life"),
    ("Total yearly cost", "=opex_total+depreciation", EUR, "OPEX + depreciation"),
    ("Labour share of yearly cost", "=SUM(labour_rows)/(opex_total+depreciation)", PCT, "labour ÷ yearly cost"),
    ("Cost per order line", "=(opex_total+depreciation)/(lines_day*days_year)", EUR2, "yearly cost ÷ lines"),
    ("Cost per order", "=(opex_total+depreciation)/(orders_day*days_year)", EUR2, "yearly cost ÷ orders"),
    ("Cost per inbound pallet", "=(opex_total+depreciation)/(inb_day*days_year)", EUR2, "yearly cost ÷ inbound pallets"),
    ("Cost per pallet position per month", "=(opex_total+depreciation)/positions/12", EUR2, "yearly cost ÷ positions ÷ 12"),
    ("Picking labour saved per year", "=(pick_today_h-pick_future_h)/7.5*paid_h_shift*days_year*labour_rate", EUR, "saved picker-h → paid h × rate"),
    ("Payback, all new (years)", "=capex_total/B12", '0.0', "CAPEX ÷ saving"),
    ("Payback, re-used racking (years)", "=capex_reuse/B12", '0.0', "CAPEX ÷ saving"),
    ("Payback, process changes only (years)", "=capex_process/B12", '0.0', "process CAPEX ÷ saving"),
    ("Loose container: in-house cost per container", "=loose_cost_year/(cont_day*days_year)", EUR, "break-even price for a contractor"),
]
for r, (t, f, fmt, txt) in enumerate(summ, 2):
    wsu.cell(row=r, column=1, value=t).font = BOLD
    c = wsu.cell(row=r, column=2, value=f); c.number_format = fmt; c.fill = TOT
    wsu.cell(row=r, column=3, value=txt)
widths(wsu, [46, 16, 40])

# ------------------------------------------------------------------ value sheets (from the analysis)
def value_sheet(title, records, cols, w):
    s_ = wb.create_sheet(title); header(s_, 1, cols)
    for r, rec in enumerate(records, 2):
        for c, k in enumerate(cols, 1):
            s_.cell(row=r, column=c, value=rec.get(k))
    widths(s_, w); s_.freeze_panes = "A2"

value_sheet("Labour", CP["labour_table"], ["Process", "Driver/day (avg)", "Unit", "Rate/h", "Peak factor", "Hours (avg)", "Hours (peak)",
                                           "FTE avg", "FTE peak", "DHL staff"], [44, 14, 10, 8, 10, 11, 11, 9, 9, 9])
value_sheet("KPIs", PC["kpis"], ["Group", "KPI", "Type", "How it is measured", "Today (data)", "Future (design)", "Target / benchmark"], [11, 30, 7, 40, 36, 48, 30])
value_sheet("Small parts", PC["small_item_options"], ["Option", "Floor m²", "Pallet positions used", "CAPEX EUR", "Pick rate lines/h",
                                                      "Picking EUR/yr", "Yearly cost EUR", "Notes"], [40, 10, 12, 12, 12, 13, 13, 80])
value_sheet("Stress test", PC["stress"], list(PC["stress"][0].keys()), [10, 20, 14, 14, 14, 26, 24])
value_sheet("Sensitivity", PC["sensitivity"], ["Driver", "Low EUR/yr", "High EUR/yr", "Swing EUR/yr"], [40, 14, 14, 14])

out = os.path.join(BASE, "deliverables", "Deliverable_3_Calculations.xlsx")
wb.save(out)
print("Saved:", out)
