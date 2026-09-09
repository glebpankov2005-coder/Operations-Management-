"""
Build the Deliverable 2 (Future State) Word document from analysis outputs.
Reads outputs/reports/phase1_metrics.json (holds capacity, capacity_plan, layout,
decision sections) + figures; writes Deliverable2_FutureState_DRAFT.docx.
"""
import os, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..")
REPORTS = os.path.join(BASE, "outputs", "reports")
FIGS = os.path.join(BASE, "outputs", "figures")
DATA = os.path.join(BASE, "data")
M = json.load(open(os.path.join(REPORTS, "phase1_metrics.json"), encoding="utf-8"))
ASSUM = json.load(open(os.path.join(DATA, "assumptions.json"), encoding="utf-8"))

RED = RGBColor(0xD4, 0x00, 0x11); DARK = RGBColor(0x22, 0x22, 0x22); GREY = RGBColor(0x66, 0x66, 0x66)
doc = Document()
normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = DARK
for lvl, sz in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    st = doc.styles[lvl]; st.font.name = "Calibri"; st.font.size = Pt(sz)
    st.font.color.rgb = RED if lvl == "Heading 1" else DARK; st.font.bold = True


def para(t, size=10.5, color=DARK, bold=False, italic=False, align=None, sa=6):
    p = doc.add_paragraph(); r = p.add_run(t)
    r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold; r.italic = italic
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(sa); return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if "**" in it:
            for i, seg in enumerate(it.split("**")):
                r = p.add_run(seg); r.bold = (i % 2 == 1); r.font.size = Pt(10.5)
        else:
            p.add_run(it).font.size = Pt(10.5)


def table(headers, rows, widths=None, arr=1):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = ""; r = t.rows[0].cells[i].paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(9.5)
            if i >= arr: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows: row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t


def figure(name, cap, w=15.5):
    p = os.path.join(FIGS, name)
    if os.path.exists(p):
        doc.add_picture(p, width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(cap, size=8.5, color=GREY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)


def footer():
    f = doc.sections[0].footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = f.add_run("Project OTM — Assa Abloy @ DHL Bemmel  ·  Deliverable 2 (Future State) DRAFT  ·  Page ")
    run.font.size = Pt(8); run.font.color.rgb = GREY
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); run._r.addnext(fld)


cp = M["capacity_plan"]; lay = M["layout"]; cap = M["capacity"]

# COVER
doc.add_paragraph().paragraph_format.space_before = Pt(55)
para("PROJECT OTM", size=13, color=RED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
para("Future-State Operational Design", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
para("Assa Abloy operation · DHL Supply Chain, Bemmel (NL_0032)", size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=26)
table(["Field", "Detail"], [
    ["Deliverable", "2 — Proposed future state for the in-scope processes"],
    ["Due date", "07-10-2026"],
    ["Concept", "Option B — Hybrid density + velocity slotting + zone/batch picking"],
    ["Envelope", "≤ 7,000 m² · ≤ 12.2 m"],
    ["Author", "________________________"],
    ["Version", "DRAFT — processes, flowcharts, layout (3D), capacity plan"],
], widths=[4.5, 11])
para("Future-state design for the recommended Option B (Deliverable 1). Workload drivers are measured "
     "from the Assa Abloy Q2-2025 data; productivities are benchmarked assumptions (Section 6) to be "
     "validated with DHL. Reproducible from analysis/*.py.", size=9, color=GREY, italic=True)
doc.add_page_break()

# 1 CONCEPT RECAP
doc.add_heading("1. Recommended concept (recap)", level=1)
para("Option B was selected via the Deliverable-1 decision matrix (robust across base and cost-driven "
     "weightings). It combines double-deep reserve racking for A/B stock, selective + cantilever racking "
     "for irregular and long goods, and a forward-pick module (carton-flow / shelving) for the case- and "
     "each-heavy demand, operated with velocity-based slotting and zone + batch/wave picking.")

# 2 FUTURE-STATE PROCESSES
doc.add_heading("2. Future-state processes", level=1)
para("Each in-scope process below states the steps, the system transaction, the physical movement, the "
     "equipment, and where the Deliverable-1 design choices are applied. Flowcharts follow.")
doc.add_heading("2.1 Receiving & putaway", level=2)
bullets([
    "**Steps:** trailer arrival → unload & RF-scan (ASN check) → quality/qty check → register stock (LODNUM) → system-directed putaway → reach-truck move → confirm.",
    "**Policy applied:** velocity-based directed putaway (A/B to double-deep reserve near pick faces; C to deep/upper); cantilever for XLONG goods (43 SKUs >240 cm).",
])
figure("flow_receiving_putaway.png", "Fig 1. Receiving & putaway flow.", 12)
doc.add_heading("2.2 Replenishment", level=2)
bullets(["**Steps:** forward face below min → WMS replen task → pull reserve pallet → top-up carton-flow/pick face → confirm.",
         "**Policy applied:** min/max replenishment; FIFO via FIF DATE where relevant."])
figure("flow_replenishment.png", "Fig 2. Replenishment flow.", 11)
doc.add_heading("2.3 Picking (zone + batch / wave)", level=2)
bullets([
    "**Steps:** orders allocated → wave release by carrier cut-off → route by profile → full-pallet / case (batched) / each pick → to consolidation.",
    "**Policy applied:** batch the 50% single-line orders; A-movers in a golden-zone forward pick; voice/RF direction; zone picking across areas.",
])
figure("flow_picking.png", "Fig 3. Picking flow (zone + batch/wave).", 12)
doc.add_heading("2.4 Consolidation, packing & shipping", level=2)
bullets([
    "**Steps:** picked totes/pallets → put-to-light sort by order → pack & print-and-apply → parcel vs pallet → stage by carrier → load & dispatch.",
    "**Policy applied:** put-to-light consolidation of multi-zone picks; pack sized to peak (labour bottleneck); parcel-heavy outbound (~18 palletised shipments/day).",
])
figure("flow_pack_ship.png", "Fig 4. Consolidation, packing & shipping flow.", 12)

# 3 POLICIES
doc.add_heading("3. Warehouse policies", level=1)
table(["Policy", "Decision", "Why (data)"], [
    ["Storage", "Hybrid: random within velocity zones; fixed forward-pick faces + random reserve", "73% single-load tail; 48.6% ABC mismatch"],
    ["Slotting", "Velocity/demand-ABC (hot/med/cold); A in golden zone; heavy low", "13% of SKUs = 80% of volume"],
    ["Zoning", "Forward-pick vs reserve; pick zones by area", "case-dominant picking (83% of lifts)"],
    ["Batching", "Batch single-line/single-unit orders; wave by carrier cut-off", "50% single-line, 24% single-unit"],
    ["Routing", "S-shape / return within aisles", "reduce travel in reserve"],
    ["Putaway", "System-directed, velocity-based to zone", "align stock to demand"],
    ["Replenishment", "Min/max forward-pick top-up", "keep pick faces filled at peak"],
    ["FIFO/FEFO", "FIFO via FIF DATE where relevant", "stock age present; most parts not date-critical (confirm)"],
], widths=[3.0, 7.5, 5.5], arr=3)

# 4 LAYOUT & ZONING (3D)
doc.add_heading("4. Layout & zoning (3-dimensional)", level=1)
para(f"The zones are sized from the capacity analysis and laid out flow-through (receiving one side, "
     f"shipping the other). Building ≈ {lay['building_L_m']:.0f} m × {lay['building_depth_m']:.0f} m, "
     f"providing ~{lay['positions_capacity']:,} pallet positions against the ~{lay['positions_target']:,} target.")
table(["Zone", "Area m²", "% of building"],
      [[z["Zone"], f"{z['Area m2']:,}", z["% of building"]] for z in lay["zones"]],
      widths=[8.5, 3.5, 3.5])
para(f"**Total floor area {lay['total_area_m2']:,} m² = {lay['pct_of_7000']:.0f}% of a 7,000 m² envelope.**",
     bold=True, sa=4)
bullets([
    "At the detailed zoning level Option B lands **slightly over 7,000 m² (107%)** once full functional "
    "areas are included — tighter than the concept-level estimate.",
    "**Mitigations if 7,000 m² is the binding Assa-only envelope (A008):** (a) raise the double-deep share "
    "or add an 8th rack level (12.2 m height allows it, subject to truck reach); (b) trim office/returns/"
    "staging; (c) switch reserve to **VNA (Option C)**, which frees ~2,000 m² and gives real headroom.",
])
figure("layout_plan.png", "Fig 5. Option B layout & zoning, to scale, with order flow (①inbound→⑥ship).", 16)
figure("layout_3d.png", "Fig 6. Option B 3D massing (isometric) — reserve racking ≈ 11.5 m within the 12.2 m limit.", 14)
figure("layout_worker_flow.png", "Fig 7. Worker movement — zone picking keeps picker paths short and separated.", 16)

# 5 CAPACITY PLAN
doc.add_heading("5. Capacity plan — labour & equipment", level=1)
m = cp["measured"]
para(f"Workload drivers (measured): outbound ≈ {m['outbound_pallets_day']:.0f} palletised shipments/day "
     f"(rest parcel), inbound ≈ {m['inbound_pallets_day']:.0f} pallets/day, forward-pick replen "
     f"≈ {m['replen_pallets_day']:.0f} pallets/day. Productivities are benchmark assumptions (A014).")
table(["Process", "Rate/h", "FTE avg", "FTE peak"],
      [[r["Process"], r["Rate/h"], f"{r['FTE avg']:.2f}", f"{r['FTE peak']:.2f}"] for r in cp["labour_table"]],
      widths=[6.5, 2.5, 2.5, 2.5])
para(f"**Total direct labour ≈ {cp['total_fte_avg']:.1f} FTE average, {cp['total_fte_peak']:.1f} FTE at peak** "
     f"(≈ {cp['total_fte_ceil_peak']} if each process is staffed separately; cross-training reduces this). "
     f"Add supervision/admin and value-add separately.", bold=True, sa=4)
figure("capacity_plan_fte.png", "Fig 8. Labour requirement by process (average vs peak).", 14)
para("Material-handling equipment (peak):", sa=4)
table(["Equipment", "Peak equip-hours/day", "Units"],
      [[r["Equipment"], f"{r['Peak equip-hours/day']:.1f}", r["Units (peak)"]] for r in cp["mhe_table"]],
      widths=[8.5, 4.0, 2.5])
para("Plus RF/voice terminals per picker, print-and-apply at pack, and a small spares allowance. Peak may "
     "need overtime or a partial second shift rather than extra permanent FTE.", size=9, color=GREY, italic=True)

# 6 ASSUMPTIONS
doc.add_heading("6. Assumptions & open questions", level=1)
para("Full register in data/assumptions.json. High-impact items for the future state:")
rows = [[a["id"], a["assumption"], a["impact"].upper()] for a in ASSUM["assumptions"] if a["impact"] == "high"]
table(["ID", "Assumption / open question", "Impact"], rows, widths=[1.5, 11.5, 2.0], arr=2)

footer()
out = os.path.join(REPORTS, "Deliverable2_FutureState_DRAFT.docx")
doc.save(out)
print("Saved:", out)
