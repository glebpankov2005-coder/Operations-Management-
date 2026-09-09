"""
Build the Word deliverable from analysis outputs.
Deliverable 1 (Analysis) - DRAFT. Reads outputs/reports/phase1_metrics.json,
data/assumptions.json and outputs/figures/*.png; writes a formatted .docx.

Regenerate any time the analysis changes:
    python analysis/build_report.py
"""
import os
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..")
REPORTS = os.path.join(BASE, "outputs", "reports")
FIGS = os.path.join(BASE, "outputs", "figures")
DATA = os.path.join(BASE, "data")

M = json.load(open(os.path.join(REPORTS, "phase1_metrics.json"), encoding="utf-8"))
M2 = M  # capacity + decision sections are written into the same metrics file by _util
ASSUM = json.load(open(os.path.join(DATA, "assumptions.json"), encoding="utf-8"))

RED = RGBColor(0xD4, 0x00, 0x11)      # DHL red accent
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK

for lvl, sz in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = RED if lvl == "Heading 1" else DARK
    st.font.bold = True


def para(text, size=10.5, color=DARK, bold=False, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold; r.italic = italic
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        # allow simple **bold:** prefix
        if "**" in it:
            parts = it.split("**")
            for i, seg in enumerate(parts):
                r = p.add_run(seg); r.bold = (i % 2 == 1); r.font.size = Pt(10.5)
        else:
            r = p.add_run(it); r.font.size = Pt(10.5)


def table(headers, rows, widths=None, align_right_from=1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v)); r.font.size = Pt(9.5)
            if i >= align_right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name, caption, width_cm=15.5):
    path = os.path.join(FIGS, name)
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = para(caption, size=8.5, color=GREY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def page_number_footer():
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Project OTM — Assa Abloy @ DHL Bemmel  ·  Deliverable 1 (Analysis) DRAFT  ·  Page ")
    run.font.size = Pt(8); run.font.color.rgb = GREY
    fld1 = OxmlElement("w:fldSimple"); fld1.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld1)


# ============================ COVER ============================
sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(60)
para("PROJECT OTM", size=13, color=RED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Warehouse Design — Analysis & Policy Pre-selection", size=22, color=DARK, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Assa Abloy operation · DHL Supply Chain, Bemmel (NL_0032)", size=12, color=GREY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

table(["Field", "Detail"], [
    ["Deliverable", "1 — Analysis of data, pre-selection & justification of policies"],
    ["Due date", "16-09-2026"],
    ["Data period analysed", "Q2 2025 (01 Apr – 30 Jun)"],
    ["Site constraints", "≤ 7,000 m² footprint · ≤ 12.2 m height"],
    ["Scope", "Receiving · Storage · Picking · Shipping (end-to-end)"],
    ["Author", "________________________"],
    ["Version", "DRAFT — analysis, design options, policies & recommendation"],
], widths=[4.5, 11])

para("This draft covers Deliverable 1 end-to-end: data analysis, storage-capacity feasibility, three "
     "design options, policy pre-selection, equipment, and a recommended concept with a decision matrix. "
     "All figures are measured from the provided Assa Abloy data or computed by the analysis scripts; "
     "nothing is fabricated, and open assumptions are listed in Section 13.",
     size=9, color=GREY, italic=True, space_after=6)
doc.add_page_break()

# ============================ 1. INTRODUCTION ============================
doc.add_heading("1. Introduction & scope", level=1)
para("DHL Supply Chain previously ran warehouse operations for four clients in Tiel and recently "
     "relocated to a new site in Bemmel under significant time pressure. The Operations Manager "
     "considers the current (post-move) design sub-optimal. The objective of this project is to "
     "develop an alternative, more efficient and effective warehouse design for the Assa Abloy "
     "operation, based on a thorough analysis of the operational data.")
para("The design covers four processes end-to-end — receiving, storage, picking and shipping — "
     "with all choices aligned across them. Facility architecture is out of scope; legal aspects "
     "(fire safety, compliance) are considered. The site envelope is a maximum footprint of "
     "7,000 m² and a maximum height of 12.2 m.")

# ============================ 2. DATA & METHOD ============================
doc.add_heading("2. Data sources & method", level=1)
inv = M["inventory"]["snapshots"]
o = M["orders"]
para("Two datasets from the Assa Abloy operation (warehouse NL_0032) form the primary source of truth:")
table(["Dataset", "Files", "Grain", "Volume"], [
    ["Storage / inventory", "3 snapshots (01-04/05/06-2025)", "Inventory detail (DTLNUM)",
     f"{inv[0]['detail_records']:,}–{inv[-1]['detail_records']:,} records"],
    ["Outbound shipping & picking", "3 months (Apr–Jun 2025)", "Order line",
     f"{o['n_order_lines']:,} lines · {o['n_orders']:,} orders"],
], widths=[3.8, 4.6, 3.6, 3.5])
para("Analysis is performed in Python (pandas). Every headline number is reproducible from the "
     "scripts in /analysis and cached in outputs/reports/phase1_metrics.json. Data-quality issues "
     "are flagged (Section 8), not silently corrected.", size=9.5, color=GREY, italic=True)

# ============================ 3. SKU PROFILE ============================
doc.add_heading("3. SKU profile", level=1)
s = M["sku"]
bullets([
    f"**{s['n_skus']:,} SKUs** in stock across **{s['n_families']} product families**.",
    f"**{s['skus_with_case_pack']:,} SKUs have a case pack; {s['skus_each_only']:,} are each-only** "
    f"(≈{round(s['skus_each_only']/s['n_skus']*100)}%) — confirming real each-pick demand.",
    f"Each size is small: median cube {s['ea_cube_m3']['median']} m³ (~500 cm³) — typical for locks, "
    f"cylinders and fittings.",
    f"Units per pallet: median **{s['units_per_pallet']['median']:.0f}** (mean skewed by high-count small parts).",
    f"Pallet weight: median **{s['pallet_weight_kg']['median']:.0f} kg** — realistic for a EURO pallet "
    f"(see data-quality flag on sentinel maxima).",
    f"Pallet-type mix (SKU level): EURO {s['asset_type_counts'].get('EURO',0):,}, "
    f"K3 {s['asset_type_counts'].get('K3',0):,}, XLONGPALLET {s['asset_type_counts'].get('XLONGPALLET',0):,} "
    f"(long goods), plus K1/K2/K5/CONS.",
    f"Hazardous SKUs: {s['hazmat_skus']} — require segregation/compliance consideration.",
])
figure("sku_pallet_weight_hist.png", "Fig 1. Pallet-weight distribution at SKU level (<p99).", 12)

# ============================ 4. INVENTORY ============================
doc.add_heading("4. Inventory & storage-position requirement", level=1)
table(["Snapshot", "Detail recs", "Loads (positions)", "Occupied locations", "Active SKUs", "Cube (m³)"],
      [[r["snapshot"], f"{r['detail_records']:,}", f"{r['loads_LODNUM']:,}",
        f"{r['occupied_locations']:,}", f"{r['active_skus']:,}", f"{r['total_cube_m3']:,.0f}"]
       for r in inv], widths=[2.7, 2.4, 3.0, 3.2, 2.2, 2.0])
ii = M["inventory"]
bullets([
    f"Loads (≈ pallet positions) grew **+{ii['loads_growth_pct_q2']}% in one quarter** "
    f"({inv[0]['loads_LODNUM']:,} → {inv[-1]['loads_LODNUM']:,}); cube +15%.",
    f"**Peak concurrent loads = {ii['peak_loads_snapshot']:,}** → baseline storage-position requirement "
    f"(before growth buffer and honeycombing losses).",
    f"Loads per SKU: mean {ii['loads_per_sku']['mean']}, median {ii['loads_per_sku']['median']:.0f}, "
    f"max {ii['loads_per_sku']['max']:.0f}; **{ii['single_load_skus_pct']}% of SKUs occupy a single load** "
    f"— a long slow-moving tail one pallet deep.",
    "Eaches held roughly flat (~2.9M) while loads rose → stock is fragmenting into more, less-full "
    "locations: a current-design inefficiency to target.",
])
figure("inventory_trend.png", "Fig 2. Inventory footprint trend, Q2 2025 (loads and occupied locations).", 13)

# ============================ 5. ORDER PROFILE ============================
doc.add_heading("5. Order & picking profile and lead time", level=1)
lpo, upl = o["lines_per_order"], o["units_per_line"]
table(["Metric", "Mean", "Median", "p95", "Max"], [
    ["Lines per order", lpo["mean"], f"{lpo['median']:.0f}", f"{lpo['p95']:.0f}", lpo["max"]],
    ["Units per line", upl["mean"], f"{upl['median']:.0f}", f"{upl['p95']:.0f}", upl["max"]],
    ["Units per order", o["units_per_order"]["mean"], f"{o['units_per_order']['median']:.0f}",
     f"{o['units_per_order']['p95']:.0f}", o["units_per_order"]["max"]],
], widths=[4.5, 2.5, 2.5, 2.5, 3.0])
lt = o["delivery_leadtime_h_arrive_to_dispatch"]
bullets([
    f"**{o['n_orders']:,} orders / {o['n_order_lines']:,} lines**; **{lpo['single_line_pct']}% single-line** "
    f"orders and **{upl['single_unit_pct']}% single-unit** lines — a small-order-heavy profile.",
    f"Pick-type split by lines: case/list **{o['pick_class_pct']['case_or_list']}%**, "
    f"each **{o['pick_class_pct']['each']}%**, full-pallet **{o['pick_class_pct']['full_pallet']}%**, "
    f"unflagged/other {o['pick_class_pct']['other']}% (see DQ flag).",
    f"Pick effort (lifts): case/list **{o['pick_method_lifts']['list_case']:,}** vs each "
    f"{o['pick_method_lifts']['trolley_each']:,} vs pallet {o['pick_method_lifts']['pallet']:,} → "
    f"**case picking is the dominant labour driver (~83% of lifts).**",
    f"Delivery lead time (order received → dispatched): **average {lt['mean']} h (~{lt['mean']/24:.1f} days)**, "
    f"median {lt['median']} h (~{lt['median']/24:.1f} days), p90 {lt['p90']} h. Right-skewed (mean > median), "
    f"so the typical order is nearer the median; measured as TRAILER DISPATCED DATE − ORDER ARRIVE DATE.",
    f"Channel mix (order type): " + ", ".join(f"{k} {v:,}" for k, v in list(o['order_type_mix'].items())[:5]) + ".",
    f"Geography: " + ", ".join(f"{k} {v:,}" for k, v in list(o['country_mix'].items())[:5]) +
    " — an NL hub serving Western Europe.",
])
figure("order_pick_class.png", "Fig 3. Order lines by pick type.", 11)

# ============================ 6. ABC ============================
doc.add_heading("6. Demand-based ABC classification", level=1)
a = M["abc"]
para("ABC is recomputed from actual outbound demand (picked quantity), not assumed. The distribution "
     "is steeper than the classic 80/20 rule:")
table(["Top X% of SKUs", "% of picked volume"],
      [[k.replace('top_', '').replace('pct_skus_do_qty_pct', '% of SKUs'), f"{v}%"]
       for k, v in a["pareto_actual"].items()], widths=[6, 6], align_right_from=1)
para("A/B/C by demand volume (A ≤ 80% cumulative, B ≤ 95%, C remainder):", space_after=4)
table(["Class", "SKUs", "% SKUs", "% volume", "% pick-lines"],
      [[c["ABC_vol"], f"{c['skus']:,}", f"{c['sku_pct']}%", f"{c['qty_pct']}%", f"{c['line_pct']}%"]
       for c in a["class_summary_volume"]], widths=[2.5, 2.5, 3, 3, 3.5])
bullets([
    "**~13% of SKUs drive 80% of volume**, yet C-class is 67% of SKUs and still 28% of picks → "
    "forward-pick / fast-lane slotting is high-value.",
    f"**Demand-ABC agrees with the WMS classification on only {a['agreement_with_wms_pct']}% of SKUs** "
    "→ strong evidence the current slotting is misaligned with real demand (re-slot).",
    f"{a['n_skus_shipped']:,} SKUs shipped in Q2; {a['skus_shipped_not_in_stock']:,} shipped were not in "
    "the latest stock snapshot (transient/cross-dock — to confirm).",
])
figure("abc_pareto.png", "Fig 4. Demand Pareto — cumulative picked quantity vs ranked SKUs.", 13)

# ============================ 7. PEAK ============================
doc.add_heading("7. Peak activity analysis", level=1)
p = M["peak"]
table(["Metric", "Average / day", "PEAK / day", "Peak ÷ Avg", "Peak date"], [
    ["Order lines", f"{p['lines']['mean']:.0f}", f"{p['lines']['peak']:.0f}", f"{p['lines']['peak_over_avg']}×", p['lines']['peak_date']],
    ["Units", f"{p['units']['mean']:.0f}", f"{p['units']['peak']:.0f}", f"{p['units']['peak_over_avg']}×", p['units']['peak_date']],
    ["Orders", f"{p['orders']['mean']:.0f}", f"{p['orders']['peak']:.0f}", f"{p['orders']['peak_over_avg']}×", p['orders']['peak_date']],
    ["Shipments", f"{p['shipments']['mean']:.0f}", f"{p['shipments']['peak']:.0f}", f"{p['shipments']['peak_over_avg']}×", p['shipments']['peak_date']],
], widths=[3.2, 3.2, 3.0, 2.6, 3.0])
bullets([
    "**Design to peak, not average.** Units peak (2.85×) is far spikier than lines (1.83×), so sizing "
    "on lines alone understates pick/pack surges.",
    f"Weekly peak {p['weekly_lines']['peak']:.0f} lines (week of {p['weekly_lines']['peak_week']}) vs "
    f"~{p['weekly_lines']['mean']:.0f} average.",
    f"Intraday peak around **{p['peak_hour']:02d}:00**; Monday–Friday only (no weekend operations), "
    "weekdays fairly even.",
])
figure("throughput_daily.png", "Fig 5. Daily shipped order-lines with average and peak.", 15)
figure("pick_hourly_profile.png", "Fig 6. Average picks per hour of day.", 13)

# ============================ 8. DATA QUALITY ============================
doc.add_heading("8. Data-quality flags", level=1)
para("Raised, not silently corrected:")
bullets([
    f"**PA WGT sentinel/impossible values** — median {s['pallet_weight_kg']['median']:.0f} kg is fine, "
    f"but p95 = {s['pallet_weight_kg']['p95']:.0f} kg and max = {s['pallet_weight_kg']['max']:.0f} kg are "
    "impossible for one pallet (9,999 looks like a placeholder). To be cleaned before rack load-rating.",
    f"**units_per_pallet extreme mean** ({s['units_per_pallet']['mean']:.0f} vs median "
    f"{s['units_per_pallet']['median']:.0f}) — use robust statistics.",
    f"**\"Other\" pick class = {o['pick_class_pct']['other']}% of lines** — no PALL/LIST/TROLLEY flag set; "
    "method needs a definition from DHL/WMS.",
    f"**Stock vs demand mismatch** — {a['skus_shipped_not_in_stock']:,} shipped-not-stocked; a further set "
    "of stocked-not-shipped SKUs are dead-stock candidates. To confirm handling (cross-dock vs obsolete).",
])

# ============================ 9. DESIGN REQUIREMENTS ============================
doc.add_heading("9. Headline design requirements (input to design options)", level=1)
bullets([
    f"**Storage:** ~{ii['peak_loads_snapshot']:,} pallet positions today, growing ~+{ii['loads_growth_pct_q2']}%/"
    "quarter → target ≈ 12,000–13,000 positions incl. buffer, within the 12.2 m / 7,000 m² envelope (A008).",
    f"**Long slow tail:** {a['class_summary_volume'][2]['sku_pct']}% C-SKUs and {ii['single_load_skus_pct']}% "
    "single-load SKUs → high-density/deep storage for the tail, fast forward-pick for A/B.",
    "**Picking is case-dominant** (~83% of lifts) with meaningful each (16% of lines) and minor full-pallet "
    "(2.6%) → mixed pick methods; case-pick productivity is the key labour lever.",
    f"**Re-slot to demand** — only {a['agreement_with_wms_pct']}% ABC agreement today.",
    f"**Capacity to peak** — {p['lines']['peak']:.0f} lines / {p['units']['peak']:.0f} units on the peak day; "
    f"~{p['peak_hour']:02d}:00 intraday peak.",
])

# ============================ 10. CAPACITY FEASIBILITY ============================
doc.add_heading("10. Storage capacity feasibility", level=1)
cap = M2["capacity"]
para(f"Design target: peak {cap['peak_loads']:,} concurrent loads × 1.15 growth/safety buffer ÷ 0.90 "
     f"target utilisation = **~{cap['design_positions_target']:,} pallet positions**. Measured pallet+load "
     f"height is low (p90 {cap['measured']['pa_hgt_p90_cm']:.0f} cm), giving a {cap['level_pitch_m']} m level "
     f"pitch and up to {cap['base_levels_height_limited']} levels under 12.2 m.")
para("Bottom-up floor area needed to hold that target, by storage concept:")
caprows = []
for c in cap["concepts"]:
    caprows.append([c["concept"].split(" (")[0], c["mhe"], str(c["levels"]),
                    f"{c['total_area_m2 (incl 35% non-storage)']:,}", f"{c['% of 7,000 m2']:.0f}%", c["fits_7000"]])
table(["Concept", "MHE", "Levels", "Total area m² *", "% of 7,000 m²", "Fits?"], caprows,
      widths=[4.0, 3.3, 1.5, 2.6, 2.4, 1.8])
para("* includes a 35% allowance for receiving, staging, pick, pack, ship, offices and circulation.",
     size=8.5, color=GREY, italic=True)
bullets([
    "**Conventional wide-aisle racking does not fit** (144% of the envelope); drive-in is also too "
    "space-hungry given the many single-load SKUs.",
    "**Double-deep (91%) and VNA (72%) fit comfortably**; narrow-aisle single-deep is borderline (107%).",
    "→ Higher-density storage is required, not optional — this shapes the options below.",
])
figure("capacity_concepts.png", "Fig 7. Floor area to hold the position target, by concept, vs the 7,000 m² envelope.", 14)

# ============================ 11. DESIGN OPTIONS ============================
doc.add_heading("11. Design options", level=1)
para("Three end-to-end options (storage × picking × MHE × policies), each aligned across the four processes.")

doc.add_heading("Option A — Conventional+ (improve the current concept)", level=2)
bullets([
    "**Storage:** selective single-deep, narrow-aisle (reach), 7 levels; two-zone ABC. ~7,500 m² (borderline).",
    "**Picking:** RF-directed discrete + simple batch; case pick from pick-face; pallet from reserve.",
    "**MHE:** reach trucks, low-level order pickers, pallet trucks, RF scanners.",
    "**Verdict:** cheapest and simplest, but only just fits, with the highest travel and labour.",
])
doc.add_heading("Option B — Hybrid density + velocity slotting + zone/batch picking  (recommended)", level=2)
bullets([
    "**Storage:** double-deep reach racking for A/B reserve + selective single-deep for irregular SKUs + a "
    "forward-pick module (carton-flow / shelving) for fast case/each; cantilever for XLONG goods. ~6,000–6,400 m².",
    "**Picking:** zone + batch/wave picking (batches the 50% single-line orders); A-movers in a golden-zone "
    "forward pick; voice/RF; put-to-light consolidation at pack.",
    "**MHE:** deep-reach trucks, order pickers, pallet trucks, voice/RF, print-and-apply.",
    "**Policies:** velocity (demand-ABC) slotting, hybrid random-within-zone storage, min/max replenishment, "
    "FIFO via FIF DATE where relevant.",
    "**Verdict:** fits with buffer, materially lower travel/labour than A, moderate CAPEX — strongest business case.",
])
doc.add_heading("Option C — High-density / semi-automated", level=2)
bullets([
    "**Storage:** VNA (man-up turret) for reserve + automated small-parts store (shuttle/AutoStore-style) for "
    "the large C each-pick tail. ~5,000 m².",
    "**Picking:** goods-to-person / pick-to-light for smalls; VNA combined storage+pick; automated sortation.",
    "**MHE:** VNA turret trucks + guidance, automation modules, minimal manual MHE.",
    "**Verdict:** densest and lowest labour, but highest CAPEX and complexity; ROI depends on labour cost.",
])

# ============================ 12. DECISION MATRIX ============================
doc.add_heading("12. Decision matrix & recommendation", level=1)
dec = M2["decision"]
wt = dec["weighted_totals"]
order = dec["ranking"]
table(["Option", "Weighted score (max 5)", "Result"],
      [[o, f"{wt[o]:.2f}", "RECOMMENDED" if o == dec["winner"] else ("Strong alternative" if wt[o] >= 3.5 else "Baseline")]
       for o in order], widths=[6.0, 4.5, 4.5])
para("Weight-sensitivity (robustness):", space_after=4)
sens = dec["sensitivity"]
table(["Option", "Base", "Cost-driven", "Space-driven"],
      [[o, f"{sens[o]['base']:.2f}", f"{sens[o]['cost_driven']:.2f}", f"{sens[o]['space_driven']:.2f}"] for o in order],
      widths=[6.0, 3.0, 3.0, 3.0])
figure("decision_matrix.png", "Fig 8. Weighted decision-matrix scores.", 12)
bullets([
    f"**Recommended: {dec['winner']}.** It wins on base and cost-driven weightings; C overtakes only when "
    "space is the dominant criterion — likely if A008 resolves to a small Assa-only envelope.",
    "**Why (data):** double-deep fits at 91% where wide-aisle (144%) does not; 83% of lifts are case picking "
    "and 50% of orders single-line → batch/zone picking is the biggest labour lever; 48.6% ABC mismatch → "
    "velocity slotting is a cheap, high-value win; 73% single-load tail → dense reserve + fast forward pick.",
    "**Conditions before locking:** confirm the 7,000 m² scope (A008); obtain labour/equipment costs (A006) to "
    "confirm B vs C in the financial model; validate double-deep against the EURO/K3/XLONG pallet mix.",
])

# ============================ 13. ASSUMPTIONS ============================
doc.add_heading("13. Assumptions & open questions", level=1)
para("Logged in data/assumptions.json. Items that materially affect the design:")
rows = [[a_["id"], a_["assumption"], a_["impact"].upper()]
        for a_ in ASSUM["assumptions"] if a_["impact"] == "high"]
table(["ID", "Assumption / open question", "Impact"], rows, widths=[1.5, 11.5, 2.0], align_right_from=2)

para("Sources: project brief (client/university) and the Assa Abloy data files. External benchmarks "
     "(productivity, equipment specs, costs, safety codes) will be added with citations as later phases use them.",
     size=9, color=GREY, italic=True)

page_number_footer()

out = os.path.join(REPORTS, "Deliverable1_Analysis_DRAFT.docx")
doc.save(out)
print("Saved:", out)
