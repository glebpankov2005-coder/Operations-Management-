"""
FINAL PORTFOLIO - one Word file for the whole group (all four students hand in the same file).
Teacher: "each student submits every portfolio entry containing the entire warehouse design, do indicate clearly which
process you are responsible for (receiving, storage, picking or shipping)".

  front matter  cover (students + their process), who owns what across D1-D3, every section per person,
                check against the assessment form and the test grid, contents
  Part 1-3      Deliverable 1, 2 and 3 exactly as generated (each with its own APA reference list)

Writes deliverables/portfolio/Portfolio_Group3_Final.docx.
Run after build_report.py, build_report_d2.py and build_report_d3.py (and docx_to_pdf.py for the PDFs).
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxcompose.composer import Composer
from team import TEAM, FULL_NAMES, PROCESS_SCOPE, OWNERS, COLOR, DOC_NAME, update_fields_on_open

BASE = os.path.join(os.path.dirname(__file__), "..")
DELIV = os.path.join(BASE, "deliverables")
OUT = os.path.join(DELIV, "portfolio"); os.makedirs(OUT, exist_ok=True)
PARTS = [("D1", "Deliverable_1_Analysis.docx", "Part 1 — Analysis",
          "Assessment form: Analysis (process requirements, suitable design choices, policies & equipment)"),
         ("D2", "Deliverable_2_Future_State.docx", "Part 2 — Future state",
          "Assessment form: Future state (process mapping, layout, zoning & capacity plan)"),
         ("D3", "Deliverable_3_Performance_Overview.docx", "Part 3 — Performance overview",
          "Assessment form: Performance overview (KPIs, capacity plan in operating costs, CAPEX & OPEX)")]
RED = RGBColor(0xD4, 0x00, 0x11); DARK = RGBColor(0x22, 0x22, 0x22); GREY = RGBColor(0x66, 0x66, 0x66)


def headings(path):
    """{number: heading text without the owner tag} of a built deliverable."""
    out = {}
    for p in Document(path).paragraphs:
        if p.style.name.startswith("Heading") and p.text[:1].isdigit():
            num = p.text.split(" ", 1)[0]
            out[num] = p.text.split("   ")[0]
    return out


HEAD = {which: headings(os.path.join(DELIV, f)) for which, f, _, _ in PARTS}


def strip_contents(doc):
    """Remove a part's own 'Contents' heading, its TOC field and the page break after it (the portfolio has one TOC)."""
    body = list(doc.paragraphs)
    for i, p in enumerate(body):
        if p.style.name == "Heading 1" and p.text.strip() == "Contents":
            drop = [p]
            for q in body[i + 1:i + 3]:
                if q._p.xpath(".//w:fldSimple") or q._p.xpath(".//w:br[@w:type='page']") or not q.text.strip():
                    drop.append(q)
            for q in drop:
                q._p.getparent().remove(q._p)
            return


class Front:
    def __init__(self, doc):
        self.doc = doc

    def para(self, t, size=10.5, color=DARK, bold=False, italic=False, align=None, sa=6):
        p = self.doc.add_paragraph()
        for i, seg in enumerate(t.split("**")):
            if seg:
                r = p.add_run(seg); r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold or i % 2 == 1; r.italic = italic
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(sa)
        return p

    def table(self, headers, rows, widths, color_col=None):
        t = self.doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = ""; r = t.rows[0].cells[i].paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                for j, seg in enumerate(str(v).split("**")):
                    if seg:
                        r = cells[i].paragraphs[0].add_run(seg); r.font.size = Pt(9.5); r.bold = j % 2 == 1
                        if color_col is not None and i == color_col and row[0] in COLOR:
                            r.font.color.rgb = COLOR[row[0]]; r.bold = True
        grid = t._tbl.tblGrid
        for i, w in enumerate(widths):
            grid.gridCol_lst[i].w = Cm(w)
            for r_ in t.rows:
                r_.cells[i].width = Cm(w)
        for ri, r_ in enumerate(t.rows):
            r_._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri < len(t.rows) - 1:                                     # keep the table (and its header) together
                for c in r_.cells:
                    for q in c.paragraphs:
                        q.paragraph_format.keep_with_next = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t


def title_of(which, key):
    h = HEAD[which]
    return h.get(key) or h.get(key + ".") or key


def build():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for lvl, sz in [("Heading 1", 15), ("Heading 2", 12.5)]:
        s = doc.styles[lvl]; s.font.name = "Calibri"; s.font.size = Pt(sz); s.font.bold = True
        s.font.color.rgb = RED if lvl == "Heading 1" else DARK
    F = Front(doc)
    # ---------------- cover: the same file for all four students
    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    F.para("PROJECT OTM — DHL ON THE MOVE", size=13, color=RED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
    F.para("Final Portfolio — Warehouse Redesign", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
    F.para("Assa Abloy operation · DHL Supply Chain, Bemmel (NL_0032)", size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)
    F.para("OPERMB04 Operations Management · HAN University of Applied Sciences · Group 3", size=10.5, color=GREY,
           align=WD_ALIGN_PARAGRAPH.CENTER, sa=20)
    F.table(["Process", "Responsible student", "Student number", "Class"],
            [[p, FULL_NAMES[n], "________________", "________________"] for p, n in TEAM.items()],
            widths=[3.2, 5.0, 4.0, 3.8], color_col=1)
    F.table(["Field", "Detail"], [
        ["Contents", "Part 1 Analysis · Part 2 Future state · Part 3 Performance overview — the whole warehouse design"],
        ["Supporting files", "Deliverable_3_Calculations.xlsx (all cost formulas); analysis/*.py, data/assumptions.json, "
                             "documentation/sources.md in the group repository"],
    ], widths=[4.0, 12.0])
    F.para("All four students hand in this same file. **Whose section is whose:** every section that belongs to one process "
           "carries the name of its owner in its heading, in the colour of that person — for example “4.2 Refilling the "
           "fast-pick area   Gleb (Storage)”. The names also appear in the contents. A section with two names was made "
           "together; a section without a name was made by the whole group. The next pages list every section per person.",
           size=9.5, color=GREY)
    doc.add_page_break()
    # ---------------- responsibilities: overview + every section per person
    doc.add_heading("Who is responsible for what", level=1)
    F.para("The project covers four processes end-to-end (project brief). Each student owns one process and wrote, calculated "
           "and checked its sections; the choices across the four processes were aligned together.", sa=4)
    rows = []
    for p, n in TEAM.items():
        rows.append([p, FULL_NAMES[n], PROCESS_SCOPE[p]] +
                    [", ".join(k.rstrip(".") for k, v in OWNERS[w].items() if p in v) for w in ("D1", "D2", "D3")])
    F.table(["Process", "Owner", "Scope", "Part 1 §", "Part 2 §", "Part 3 §"], rows, widths=[2.2, 3.0, 4.6, 1.8, 2.6, 1.8], color_col=1)
    F.para("**Whole group (no name in the heading):** data preparation and quality checks, the operations-strategy analysis, "
           "design requirements, design options and the decision matrix (Part 1); concept, policies overview, layout and zoning, "
           "labour plan, machinery and the control loop (Part 2); summary, KPI framework, zone guide, stress test, costs, payback "
           "and sensitivity (Part 3).", size=9.5)
    for proc, name in TEAM.items():
        h = doc.add_heading(f"Sections of {FULL_NAMES[name]} — {proc}", level=2)
        for r_ in h.runs:
            r_.font.color.rgb = COLOR[proc]
        mine = []
        for w in ("D1", "D2", "D3"):
            for k, v in OWNERS[w].items():
                if proc in v:
                    others = [f"{TEAM[o]} ({o})" for o in v if o != proc]
                    mine.append([f"Part {w[1]} ({DOC_NAME[w].split(' — ')[1]})", title_of(w, k.rstrip(".")),
                                 "with " + ", ".join(others) if others else "own section"])
        F.table(["Part", "Section", "Made"], mine, widths=[4.4, 8.6, 3.0])
    doc.add_page_break()
    # ---------------- assessment check
    doc.add_heading("How this portfolio meets the assessment form", level=1)
    F.para("Knock-out criteria: **APA referencing** — every part cites its sources as (author, year) in the text and ends with an "
           "APA 7 reference list; conversations with the teacher and DHL are cited as personal communications. **Hand-in "
           "deadline** — handed in via Brightspace before the deadline.", sa=4)
    F.table(["Deliverable (weight)", "Standard for ‘Excellent’", "Where to find it"], [
        ["Analysis (1/3)", "Requirements based on an excellent analysis of the data; design aspects, policies and equipment all "
         "suitable, substantiated with appropriate sources",
         "Part 1 §3–10 data analysis per process (SKUs, inventory, inbound, orders, outbound, ABC, peaks, data quality); §11 operations "
         "strategy; §12 requirements per process; §13 capacity; §14–15 options and decision matrix · Part 2 §3 policies, §6.3 "
         "equipment · Part 3 §6.1 machinery check"],
        ["Future state (1/3)", "Process fully in line with customer requirements and design choices; detailed elaboration; "
         "excellent input for the new operation",
         "Part 2 §2 process maps (swimlane + flowcharts), §3 policies, §4 docks and departure plan, §5 layout and zoning (2D, 3D), "
         "§6 capacity plan (labour, equipment), §7 planning and control · Part 3 §3–5 zones, storage and sorting explained"],
        ["Performance overview (1/3)", "Relevant KPIs; detailed capacity plan with expected operating costs; aligned with the "
         "analysis and future state",
         "Part 3 §2 own + lecture KPIs with expected values, each linked to a performance objective and an owner; §2.4 stress "
         "test; §7 CAPEX, OPEX, cost per unit, payback, sensitivity; calculation workbook"],
    ], widths=[3.0, 5.6, 7.4])
    F.para("Test grid (module learning outcomes TWM20/TWM21):", sa=4)
    F.table(["Learning outcome (weight)", "Assessment criteria", "Where to find it"], [
        ["Alignment of operations performance and business strategy (20%)",
         "categorisation of operations processes; classification of process performance; relationship operations–strategy; "
         "evaluation of current alignment",
         "Part 1 §11.1 the four processes and their type; §11.2 the four V's; §11.3 performance objectives, order winners and "
         "qualifiers; §11.4 current alignment and gaps · Part 3 §2 KPIs per objective"],
        ["Designs a process given product/service and performance characteristics (40%)",
         "classification of expected process performance; appropriateness of choices; quality of process mapping",
         "Part 3 §2 expected performance per KPI; Part 1 §12–15 and Part 2 §2–6 choices per process from the data; Part 2 §2 "
         "swimlane and flowcharts; Part 3 §3–5 zone, storage and put-wall figures"],
        ["Proposes methods for planning and control (40%)",
         "relevance of the factors considered; quality of planning and control decisions",
         "Part 2 §7 loading, sequencing, scheduling and control per process with the factors from the data; §4.5 departure plan "
         "and waves; §6 capacity plan · Part 3 §2 KPIs as control measures, §2.4 stress test"],
    ], widths=[3.6, 5.0, 7.4])
    doc.add_page_break()
    # ---------------- contents
    F.para("Contents", size=15, color=RED, bold=True, sa=6)            # not a heading: the TOC does not list itself
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "(In Word: right-click → Update Field to build the table of contents.)"
    r.append(t); fld.append(r); p._p.append(fld)
    # footer
    f = doc.sections[0].footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = f.add_run("Project OTM — Group 3 final portfolio  ·  " +
                    "  ·  ".join(f"{n} ({p})" for p, n in TEAM.items()) + "  ·  Page ")
    run.font.size = Pt(8); run.font.color.rgb = GREY
    pf = OxmlElement("w:fldSimple"); pf.set(qn("w:instr"), "PAGE"); run._r.addnext(pf)
    # ---------------- parts
    comp = Composer(doc)
    for which, fname, title, form in PARTS:
        comp.doc.add_page_break()
        comp.doc.add_heading(title, level=1)
        q = comp.doc.add_paragraph(); rr = q.add_run(form); rr.italic = True; rr.font.size = Pt(9.5); rr.font.color.rgb = GREY
        part = Document(os.path.join(DELIV, fname))
        strip_contents(part)
        comp.append(part)
    update_fields_on_open(comp.doc)
    out = os.path.join(OUT, "Portfolio_Group3_Final.docx")
    comp.save(out)
    return out


if __name__ == "__main__":
    print("Saved", build())
